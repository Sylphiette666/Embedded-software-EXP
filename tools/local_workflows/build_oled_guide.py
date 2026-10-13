"""Create the lab 8 OLED supplement. Uses bundled python-docx; no hardware access.

The caller must run the document-authoring marker before invoking this builder.
Evidence notes are supplied explicitly using --validation-text. This document
never upgrades compilation or emulation into an on-board test result.
"""
from __future__ import annotations

import argparse
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(r"D:\MDK5\F103_Labs")
OUTPUT = Path(r"D:\Work\嵌入式实验\实验八_OLED接线与调试流程_20261013.docx")


class Guide:
    def __init__(self):
        self.doc = Document()
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.top_margin, sec.bottom_margin = Cm(1.9), Cm(1.8)
        sec.left_margin, sec.right_margin = Cm(2), Cm(2)
        sec.header_distance, sec.footer_distance = Cm(0.75), Cm(0.75)
        for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Header", "Footer"):
            style = self.doc.styles[name]
            style.font.name = "Calibri"
            style.font.color.rgb = RGBColor(0, 0, 0)
            style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
            # The bundled blank document may carry a blue Title paragraph rule.
            # Strip all paragraph borders from these prose styles explicitly.
            ppr = style.element.get_or_add_pPr()
            for border in list(ppr.findall(qn("w:pBdr"))):
                ppr.remove(border)
        normal = self.doc.styles["Normal"]
        normal.font.size = Pt(10.5)
        normal.paragraph_format.line_spacing = 1.12
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.widow_control = True
        for name, size in (("Title", 21), ("Heading 1", 16), ("Heading 2", 12)):
            st = self.doc.styles[name]
            st.font.size = Pt(size)
            st.font.bold = True
            st.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "黑体")
            st.paragraph_format.space_before = Pt(9)
            st.paragraph_format.space_after = Pt(7)
            st.paragraph_format.keep_with_next = True
        self.doc.styles["Title"].paragraph_format.space_before = Pt(0)
        footer = sec.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = footer.add_run("实验八 OLED 补充指南  |  ")
        r.font.size = Pt(9)
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        footer._p.append(field)
        props = self.doc.core_properties
        props.title = "实验八 OLED 接线与调试流程"
        props.subject = "普中精灵 STM32F103C8T6 开发板 SSD1306 I2C OLED 实验指导"
        props.author = ""
        props.comments = ""

    def p(self, text, bold=False, size=None):
        p = self.doc.add_paragraph()
        r = p.add_run(text)
        r.bold = bold
        if size:
            r.font.size = Pt(size)
        return p

    def h(self, title):
        return self.doc.add_paragraph(title, "Heading 2")

    def page(self, title, first=False):
        p = self.doc.add_paragraph(title, "Title" if first else "Heading 1")
        if not first:
            p.paragraph_format.page_break_before = True
        return p

    def steps(self, values):
        for i, value in enumerate(values, 1):
            p = self.p(f"{i}  {value}")
            p.paragraph_format.left_indent = Cm(0.48)
            p.paragraph_format.first_line_indent = Cm(-0.48)

    def table(self, headers, rows, widths):
        t = self.doc.add_table(rows=1, cols=len(headers))
        t.autofit = False
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pr = t._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            edge = OxmlElement(f"w:{side}")
            for key, value in (("val", "single"), ("sz", "5"), ("color", "D9D9D9")):
                edge.set(qn(f"w:{key}"), value)
            borders.append(edge)
        pr.append(borders)
        for c, width in zip(t.columns, widths):
            c.width = Cm(width)
        for values, row in [(headers, t.rows[0])] + [(values, t.add_row()) for values in rows]:
            for j, (cell, value) in enumerate(zip(row.cells, values)):
                cell.width = Cm(widths[j])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcpr = cell._tc.get_or_add_tcPr()
                margins = OxmlElement("w:tcMar")
                for side, val in (("top", 80), ("bottom", 80), ("left", 100), ("right", 100)):
                    item = OxmlElement(f"w:{side}")
                    item.set(qn("w:w"), str(val))
                    item.set(qn("w:type"), "dxa")
                    margins.append(item)
                tcpr.append(margins)
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.03
                run = p.add_run(str(value))
                run.font.size = Pt(10)
                if row is t.rows[0]:
                    run.bold = True
            trpr = row._tr.get_or_add_trPr()
            cant = OxmlElement("w:cantSplit")
            trpr.append(cant)
        for cell in t.rows[0].cells:
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "E7EDF3")
            cell._tc.get_or_add_tcPr().append(shade)
            for run in cell.paragraphs[0].runs:
                run.bold = True
        repeat = OxmlElement("w:tblHeader")
        t.rows[0]._tr.get_or_add_trPr().append(repeat)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)
        return t

    def code(self, text):
        p = self.p(text, size=9)
        p.paragraph_format.space_after = Pt(5)
        for r in p.runs:
            r.font.name = "Consolas"
        return p


def build(output: Path, validation: str):
    g = Guide()
    g.page("实验八 OLED 接线与调试流程", first=True)
    g.p("普中精灵 STM32F103C8T6 完整开发板   2026 年 10 月 13 日", size=10)
    g.p("实验八改用 OLED 显示完成文字。K1 控制 D1 闪烁，K4 使屏幕显示“嵌入式实验完成啦”。本指南补充实验八的接线、下载、调试和验收步骤；实验一至七沿用原指南。")
    g.p("当前驱动按 SSD1306 控制器、128×64 分辨率编写。已确认的是四针 I²C 接口；仅凭 GND、VCC、SCL、SDA 不能确定控制器和分辨率。若模块是 SH1106 或 128×32，须按实际型号调整驱动后再验收。", bold=True)
    g.h("先连接 OLED 四根线")
    g.table(["OLED 丝印", "开发板连接点", "作用"], [
        ("GND", "GND", "公共地"),
        ("VCC", "3V3 或 3.3V", "OLED 模块供电"),
        ("SCL 或 SCK", "PB10", "I²C 时钟线"),
        ("SDA", "PB11", "I²C 数据线"),
    ], [4.0, 5.0, 8.0])
    g.p("从核心板标有 PB10、PB11 的 GPIO 引脚引出。以上表格表示连接关系，不表示模块针脚从左到右的物理顺序。逐个读取丝印；底板标有 OLED 的插座也要先确认其实际连线，不能直接按名称插入。")
    g.steps([
        "先拔下底板 USB 和 ST Link USB，再连接上述四根线；核心板继续插在底板上，无须取下。",
        "确认 VCC 接 3.3 V、GND 共地。SCL 和 SDA 需要上拉到 3.3 V，常见模块已带上拉电阻；不要把信号线通过上拉接到 5 V。",
        "BOOT0/B0 与 BOOT1/B1 均放 GND 侧。保留底板原有 LED、按键跳线；D1 对应 PA0，K1 对应 PA15，K4 对应 PA12。",
        "重新连接底板 CH340 附近的 USB 数据口供电。该口同时提供串口；不要把核心板自身 USB 口当作 CH340 串口。",
    ])
    g.h("保留下载器的三根信号线")
    g.p("ST Link 的 SWDIO/DIO 接 PA13/DIO，SWCLK/CLK 接 PA14/CLK，GND 接 GND。下载器若有专用 VTref/VAPP 参考输入，按其说明接目标 3.3 V；不要把 3.3 V/5 V 电源输出脚并接到已由底板 USB 供电的板上。K2、K3 占用 SWD 引脚，调试时不用它们。")

    g.page("打开工程并下载实验八")
    g.h("本机软件入口")
    g.table(["软件", "用途与路径"], [
        ("Keil µVision", "打开工程、编译；SWD 连通后下载及断点调试\nD:\\Keil5\\UV4\\UV4.exe"),
        ("FlyMCU", "通过底板 CH340 执行串口 ISP 下载\nD:\\FlyMCU\\FlyMcu.exe"),
        ("Tera Term", "查看辅助串口日志，115200、8N1、无流控、UTF 8\nD:\\TeraTerm\\ttermpro.exe"),
    ], [3.0, 14.0])
    g.p("现有编译环境为 Arm Compiler 6.24 和 STM32F1xx DFP 2.4.1；目标芯片为 STM32F103C8T6，Flash 64 KB、RAM 20 KB。", size=10)
    g.h("选择正确目标并重新编译")
    g.steps([
        "打开 Keil，选择 Project → Open Project，打开下列工程文件。目标下拉框选择 08_Application。",
        "执行 Project → Rebuild all target files。等待编译结束，检查本次输出中错误和警告的数量。不要把旧的成功日志当作本次结果。",
        "编译产物应为下列 HEX。检查文件的修改时间与刚完成的编译一致，再选择下载方式。",
    ])
    g.code(r"D:\MDK5\F103_Labs\MDK\F103_Labs.uvprojx")
    g.code(r"D:\MDK5\F103_Labs\Build\Keil_08_Application\08_Application.hex")
    g.h("使用串口 ISP 下载")
    g.steps([
        "在设备管理器展开“端口 COM 和 LPT”，找到 CH340 的实际 COM 号。历史记录为 COM4，重新插拔后必须核对。退出 Tera Term 等占用此端口的软件。",
        "打开 FlyMCU，选择 STM32 串口下载、实际 COM 口和本页列出的 HEX，启用写入后校验。手动进入 ISP 时将 BOOT0 置 1、BOOT1 置 0，然后按 RST。",
        "开始下载，必要时按软件提示再次复位。只写入应用程序，不修改读保护或选项字节。板载自动下载电路的脚本方式见原全实验指南，不把手动 BOOT 步骤与脚本自动握手混用。",
        "确认下载和校验均成功，关闭下载端口；将 BOOT0 恢复 0、BOOT1 保持 0，再按 RST 运行。打开 Tera Term 后可再复位一次保存启动日志。",
    ])
    g.h("SWD 连通后才使用 Keil 下载和调试")
    g.p("2026 年 10 月 4 日的记录中，独立 ST Link 能被电脑识别，但尚不能连接目标 MCU。先在 Options for Target → Debug 选择 ST Link Debugger，在 Settings 中选择 SW 接口及 1 MHz，并确认能识别目标芯片；成功后再执行 Flash → Download。仍提示无法连接时，采用已用过的 ISP 路径继续验证 OLED，SWD 子项保留待排查。")

    g.page("观察 OLED 并逐项验证")
    g.p("完成下载并恢复 BOOT0=0 后，按 RST 从初始状态开始。必须直接观察实体屏幕及 D1；PC 上的中文日志只用于帮助定位程序执行位置。")
    g.h("复位后的预期画面")
    g.table(["屏幕位置", "预期内容"], [
        ("第一行", "LAB8 OLED READY"),
        ("第二行", "BLINK: OFF"),
        ("第三行", "初始为空；按 K4 后显示 嵌入式实验完成啦"),
        ("第四行", "K1:BLINK K4:MSG"),
    ], [4.0, 13.0])
    g.p("驱动依次探测七位地址 0x3C、0x3D。它们是七位地址；不要把写地址字节 0x78、0x7A 当成新的七位地址。屏幕会分块刷新，按键后允许一个短暂的刷新过程。")
    g.h("按顺序操作")
    g.steps([
        "复位后确认 D1 熄灭，屏幕显示初始内容。若屏幕不亮，先按第 4 页排查，不要把串口正常当作屏幕通过。",
        "短按 K1 后松开。D1 应立即点亮，OLED 第二行变为 BLINK: ON，随后 D1 每 1 秒翻转一次。完整亮灭周期约 2 秒。",
        "连续观察至少 3 次翻转，在闪烁期间短按 K4。第三行应出现完整的“嵌入式实验完成啦”，D1 继续闪烁。",
        "长按 K4 至少 1 秒，再松开。业务事件只应执行一次；因为文字保留在屏幕上，单看同一行文字无法判断重复次数，可检查辅助串口是否只新增一行中文日志。",
        "再次按 K1，OLED 显示 BLINK: OFF，D1 立即熄灭并保持。已显示的中文应继续保留。",
        "在停止闪烁状态下再次短按 K4，中文仍应完整显示；再按 K1 启动，中文与闪烁状态行应同时保留。",
        "重复启停和显示至少 3 轮。按住任一键不应连续切换；两键同时按下时，应同时处理启停与显示两个事件。",
        "保存能够看清文字的屏幕照片，以及包含 K1、K4 操作和 D1 连续闪烁的短视频。记录下载文件、操作时间、现象和不符合项。",
    ])
    g.h("验证未接屏时按键仍能工作")
    g.p("这是可选故障检查。先断电，再拔开 OLED 四根线，重新上电；程序应仍能响应 K1 并控制 D1 闪烁。检查完先断电恢复接线，再复位。显示初始化失败或运行中发生总线错误后，本版停止 OLED 传输；修线后须复位重新初始化。不要带电拔插模拟故障。")

    g.page("使用断点定位问题")
    g.p("以下操作需要 SWD 已经连通。进入 Debug → Start/Stop Debug Session 后，在主程序初始化之后设置断点。先看状态，再清除断点全速运行；单步和暂停会影响闪烁时间，不能用暂停过程测量 1 秒间隔。")
    g.h("观察这些变量")
    g.table(["变量", "应如何解释"], [
        ("g_oled_present", "为 1 表示驱动初始化成功；为 0 表示离线。它不能证明像素已经在实体屏幕正确显示。"),
        ("g_oled_address", "初始化成功时应为探测到的七位地址 0x3C 或 0x3D。"),
        ("g_oled_error_count", "总线错误累计计数；失败后检查接线、上拉和模块型号。"),
        ("g_oled_frame_count", "完整显示帧发送计数。状态或文字改变后应完成刷新；画面未变时不要求持续增加。"),
        ("g_oled_framebuffer", "1024 字节显存，按 8 像素高的页排列。只证明待发送的像素内容。"),
        ("g_millis 与 IRQ 计数", "毫秒节拍继续增加；K1、K4 分别对应 g_key0_irq_count、g_key1_irq_count。原始中断可受机械抖动影响，业务按键已消抖。"),
    ], [5.2, 11.8])
    g.p("主循环先处理按键事件和 LED 定时，再调用 OLED_Task 分块发送显示数据。OLED 传输放在主循环，外部中断处理函数只记录事件。检查流程时可依次看 OLED_Init、OLED_SetBlink、OLED_ShowCompleted 和 OLED_Task。")
    g.h("常见现象与检查顺序")
    g.table(["现象", "先检查"], [
        ("完全黑屏", "VCC 是否为 3.3 V、GND 是否共地；SCL/SDA 是否接反；PB10/PB11 是否接对；初始化是否成功。修线后复位。"),
        ("有地址应答但无字", "核对控制器和分辨率。SH1106、128×32 等不能仅靠四针接法确认兼容；仍需实体画面检查。"),
        ("乱码或画面偏移", "确认 SSD1306 128×64；核对供电、接触和上拉。先保留较短连线，再核对驱动与实际屏幕方向。"),
        ("屏幕停止更新", "查看错误计数和 ready 状态；断电检查连线，恢复后复位。本版不会无限循环等待故障总线。"),
        ("K1 或 K4 没响应", "核对实际键为 K1/PA15、K4/PA12；不要用 RST、K2、K3；确认程序在全速运行，未停在断点。"),
        ("PC 有中文但屏黑", "只说明辅助串口路径执行。OLED 验收仍未通过，继续检查供电、I²C 应答和控制器型号。"),
    ], [4.0, 13.0])

    g.page("源码入口与验收记录")
    g.h("当前工程文件")
    g.p("工程根目录为 D:\\MDK5\\F103_Labs。OLED 源码只加入 08_Application；实验一至七仍使用原来的显示与按键流程。完整源码包为“普中精灵F103_完整源码_OLED版_20261013.zip”，以下文件可在 Keil 或文本编辑器中打开。")
    g.table(["根目录下的文件", "作用"], [
        (r"Labs\08_Application\main.c", "实验八初始化、按键事件、LED 定时和显示任务调度"),
        (r"Labs\08_Application\oled.c", "PB10/PB11 软件 I²C、SSD1306 命令与分块刷新"),
        (r"Labs\08_Application\oled.h", "OLED 接口函数与调试变量声明"),
        (r"Labs\08_Application\oled_font.h", "英文字符和完成语句的字模数据"),
        (r"App\platform.c", "公共时钟、LED、按键消抖、EXTI 和辅助串口"),
        (r"MDK\F103_Labs.uvprojx", "八个编译目标和实验八 OLED 源文件配置"),
    ], [8.0, 9.0])
    g.p("原 PC 显示版主程序保存在工程的 Backup\\20261013_lab8_PC_original_main.c，便于对照修改。", size=10)
    g.h("本次验证状态")
    g.p(validation)
    g.p("OLED 尚未在实体板上完成下载和屏幕观察验收。2026 年 10 月 4 日的旧记录只覆盖实验 1、2、3、4、6、7 的部分下载或串口检查；实验 5、8 当时未实测。旧记录不能作为本版 OLED 显示、机械按键或 D1 光学现象的通过证据。", bold=True)
    g.h("上板后填写")
    g.table(["记录项目", "实测填写"], [
        ("屏幕型号和分辨率", "控制器 ________   分辨率 ________   七位地址 ________"),
        ("下载与观察", "日期 ________   下载方式 ________   HEX 校验 ________"),
        ("复位和 K1", "初屏 ________   启停 ________   1 秒翻转 ________"),
        ("K4 与联合操作", "完整中文 ________   闪烁不中断 ________   长按 ________"),
        ("证据与结论", "照片或视频路径 ______________________________\n未通过项及处理 ______________________________"),
    ], [4.0, 13.0])
    g.p("验收依据是当前固件、实际连线和实体可见现象。源码编译成功、离线模型通过、串口日志正确分别记录，只有完成板上观察后才能填写 OLED 实测通过。")
    g.doc.save(output)
    print(output)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=OUTPUT)
    parser.add_argument("--validation-text", type=Path, required=True)
    args = parser.parse_args()
    build(args.output, args.validation_text.read_text(encoding="utf-8-sig").strip())
