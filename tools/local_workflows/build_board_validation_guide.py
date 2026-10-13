"""Build the F103 on-board verification guide from explicit evidence.

Run with the Codex bundled Python. This module does not access serial ports,
debug probes, or target memory. Merely importing it does not create a DOCX.
The document authoring marker must be run by the caller before --output.
"""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

PROJECT = Path(r"D:\MDK5\F103_Labs")
SUMMARY = PROJECT / "Evidence/20261004/validation_summary.json"
OUTPUT = Path(r"D:\Work\嵌入式实验\普中精灵F103_上板验证与源码调试流程_20261004.docx")

STATUSES = {
    "passed": "通过", "pass": "通过", "board_pass": "通过", "partial": "部分验证",
    "failed": "未通过", "fail": "未通过", "not_tested": "待测", "pending": "待测",
    "blocked": "待补测", "not_applicable": "不适用", "skipped": "未测",
}
METHODS = {
    "physical_button": "实际按键操作", "swd_register": "板上 SWD 寄存器读回",
    "swd_injection": "板上 SWD 注入条件", "serial": "实际串口收发",
    "instrument": "外部仪器测量", "visual": "实物观察", "offline": "离线检查",
    "isp": "串口 ISP 下载及校验", "swd": "ST Link SWD 下载及校验",
    "isp_readback": "ROM ISP 全片读回校验", "serial_report": "实际串口中的固件报告",
    "gpio_injection": "GPIO 电平注入", "user_observation": "用户实物观察",
}

LABS = [
    {
        "target": "01_Platform", "title": "实验一 开发平台与两种下载方式",
        "purpose": "建立编译 下载 运行 调试的完整流程，并分别验证串口 ISP 和 ST Link 下载。",
        "steps": [
            "打开 F103_Labs.uvprojx，选择 01_Platform。执行 Project → Rebuild all target files，确认当前编译输出，记录错误与警告数量。",
            "按公共下载流程通过 CH340 串口 ISP 写入本目标 HEX，并执行写入校验。将 BOOT0 恢复 0 后复位，记录串口启动信息和 D1 状态。",
            "再通过 ST Link 的 SWD 接口下载同一目标，并保存写入校验成功记录。若 SWD 尚未连通，可用已验证的 ISP 继续后续实验，但实验一的 SWD 子项仍须保留未通过或待补测。",
            "用 Tera Term 打开实验板 COM 口，复位后确认 LAB 1 和 HCLK=72000000。发送英文问号 ? 查询 STATE，观察 D1 每 500 ms 翻转。",
            "进入 Debug 会话，在 main 初始化后设断点。查看 SystemCoreClock 与 g_millis；继续运行再暂停，确认毫秒计数增加，然后清除断点全速运行。",
        ],
        "debug": "从 Reset_Handler 跟踪到 SystemInit 和 main；区分启动文件、系统时钟、标准库和 Labs 主程序。串口打开后再复位才能完整保存启动行。",
        "accept": "ISP 和 ST Link 均完成下载及校验；板上程序能够启动、闪烁并回复串口查询。任一种下载方式未实操时，对应项保留待测。",
    },
    {
        "target": "02_Register", "title": "实验二 直接使用寄存器控制 LED",
        "purpose": "用 RCC GPIO 寄存器控制底板 D1 和 D2 交替点亮。",
        "steps": [
            "选择 02_Register，重新编译并下载本目标 HEX；打开串口后复位，确认 LAB 2。",
            "观察底板前两盏 LED：D1 亮 D2 灭约 500 ms，然后 D1 灭 D2 亮约 500 ms，连续观察至少五轮。",
            "在 Register_LED_Init 后设断点。确认 RCC APB2ENR 的 IOPAEN 已置位，GPIOA CRL 低 8 位为 0x22，即 PA0 PA1 为 2 MHz 推挽输出。",
            "在两次 BSRR 写入之后分别暂停。第一种状态 GPIOA ODR 低两位应为 10，第二种应为 01；D1 D2 为低电平点亮。",
            "清除断点，检查串口 LED0=1 LED1=0 与 LED0=0 LED1=1 交替出现，并与实体灯对应。LED0=1 表示亮，并非 PA0 为高电平。",
        ],
        "debug": "本实验的 LED 初始化和切换直接访问寄存器。Time_Init、Delay_Ms、UART_Init 只提供时间与日志辅助，不改变寄存器点灯的实验目的。",
        "accept": "两盏 LED 交替亮灭，GPIO 配置和 ODR 读回符合代码。串口状态只证明程序执行；LED 外观应另有实物观察记录。",
    },
    {
        "target": "03_StdPeriph", "title": "实验三 使用标准外设库控制 LED",
        "purpose": "通过 STM32F1 标准外设库使 D1 亮一秒 灭一秒。",
        "steps": [
            "选择 03_StdPeriph，编译并下载，复位后确认 LAB 3。D1 应亮 1 s、灭 1 s，完整亮灭周期约 2 s。",
            "连续观察至少五个完整周期；保存 ON 和 OFF 交替的串口记录，核对日志与实体 LED 同步。",
            "在 App/platform.c 的 LED_Init 中逐行查看 RCC_APB2PeriphClockCmd、GPIO_StructInit 和 GPIO_Init，确认 PA0 PA1 的输出配置。",
            "在 LED_Write 中查看参数 index 与 on。on=1 时调用 GPIO_WriteBit 输出 Bit_RESET，适配底板 LED 的低电平有效电路。",
            "清除所有断点后重新计时。断点会暂停应用，不能把调试暂停造成的延时算入正常闪烁周期。",
        ],
        "debug": "比较实验二的直接寄存器写入与本实验的库函数调用。Delay_Ms 以 g_millis 的无符号时间差等待，SysTick_Handler 每次增加 1。",
        "accept": "标准库初始化正常，D1 稳定亮 1 s 灭 1 s。若只检查了日志或 ODR，应明确记录观测方式。",
    },
    {
        "target": "04_Keys", "title": "实验四 按键轮询与消抖",
        "purpose": "用 K1 同时反转 D1 D2，用 K4 产生互斥交替状态。",
        "steps": [
            "选择 04_Keys，编译下载并复位，两灯初始熄灭。程序 KEY0 对应 K1，KEY1 对应 K4；不要使用 K2 K3 或 RST 代替。",
            "短按 K1 后松开，两灯应同时亮；再次按 K1，两灯同时灭。每次松开稳定超过 20 ms 后再按。",
            "复位后短按 K4，D1 亮 D2 灭；再次短按 K4，D1 灭 D2 亮。连续重复，验证 K4 独立保存互斥状态。",
            "在互斥状态下按 K1，两个 LED 各自反转。随后按 K4，最终状态由 K4 的独立互斥变量决定。",
            "长按每个键至少 1 s，应只发生一次业务动作；同时按下两键时，程序先处理 K1，再处理 K4。保存 STATE 序列。",
        ],
        "debug": "Keys_Poll 要求输入稳定 20 ms 才确认按下。检查 PA15 PA12 未按为 1、按下为 0；检查 AFIO_MAPR 已关闭 JTAG 并保留 SWD。软件注入事件只能验证状态机，不能替代实体按键与电路消抖测试。",
        "accept": "K1 同时反转两灯，K4 依次产生 10 和 01；每次稳定按压只触发一次。记录实际按键与注入验证的区别。",
    },
    {
        "target": "05_EXTI", "title": "实验五 外部中断控制 LED",
        "purpose": "用 K1 的 EXTI15 翻转 D1，用 K4 的 EXTI12 翻转 D2。",
        "steps": [
            "选择 05_EXTI，编译下载并复位，两灯初始熄灭。按 K1 K4 K1 K4 的顺序各一次，预期 D1D2 状态为 10、11、01、00。",
            "在 EXTI15_10_IRQHandler 中设断点。分别按 K1 和 K4，确认都进入该共享处理函数，再检查是哪条线挂起。",
            "查看 EXTI IMR 与 FTSR 的第 15 位和第 12 位是否启用。两个独立 if 分别处理各中断源并清除各自挂起位。",
            "清除断点，长按每键一次应只触发一次 LED 业务动作。同时按下两键，检查两灯均能翻转，不能遗漏共享 IRQ 的第二个源。",
            "保存 STATE 的 IRQ0 IRQ1 与灯状态。机械抖动可能让原始 IRQ 计数增加多次；消抖后的业务动作仍应只执行一次。",
        ],
        "debug": "中断只记录计数和锁存事件，主循环执行消抖并更新 LED。向 EXTI 软件触发寄存器写值属于中断注入测试，不能据此声称机械按键产生的真实下降沿已验证。",
        "accept": "两实体按键通过中断分别控制 LED，共享 IRQ 不遗漏两个挂起源；实测记录须说明中断来自按键还是调试注入。",
    },
    {
        "target": "06_USART", "title": "实验六 串口命令控制 LED",
        "purpose": "通过 PC 发送大写 Y 点亮 D1，发送大写 N 熄灭 D1。",
        "steps": [
            "选择 06_USART，编译下载；关闭 FlyMCU 等占用 COM 口的软件。Tera Term 设置为 115200、8N1、无流控、UTF 8。",
            "复位板卡，确认收到 LAB 6 和 HCLK=72000000。若没有启动行，检查底板 USB1 与 P4 的 A9–UR、A10–UT 跳帽。",
            "以文本方式发送大写 Y，预期回复 OK Y LED0=1 且 D1 点亮。再发送大写 N，预期回复 OK N LED0=0 且 D1 熄灭。",
            "在开灯状态发送小写 y、X 和回车，再查询 ?，灯状态应保持；在关灯状态重复无效字符测试。Y 的 ASCII 值为 0x59，N 为 0x4E。",
            "需要跟踪接收时在 USART1_IRQHandler 设断点，查看 DR 中的字节，再继续至主循环的 command 判断。完成后清除断点复测。",
        ],
        "debug": "PA9 为 USART1 TX，PA10 为 RX。72 MHz 外设时钟和 115200 波特率时 BRR 应为 0x271。调试暂停可能导致接收溢出，应查看 RX_ERRORS 并在全速运行状态重新测试。",
        "accept": "PC 经实际 CH340 发送 Y N，固件回复及 ODR 状态符合预期；普通无效字符不改变报告状态。目标引脚电平和实体 LED 现象需另有实测或观察证据。",
    },
    {
        "target": "07_SysTick", "title": "实验七 SysTick 毫秒定时",
        "purpose": "配置 1 ms SysTick 节拍，使 D1 每 1000 ms 翻转一次。",
        "steps": [
            "选择 07_SysTick，编译下载并复位。确认串口 HCLK=72000000；若出现时钟警告，先排查 8 MHz 外部晶振。",
            "检查 SystemCoreClock=72000000，SysTick LOAD=71999，CTRL 低三位为 111。SysTick_Handler 每次令 g_millis 加 1。",
            "全速运行，连续保存至少十次 STATE，检查相邻 ms 值约增加 1000；同时检查 D1 每秒翻转，完整亮灭周期约 2 s。",
            "理解 main 中 now-last>=1000 与 last+=1000 的周期安排。计时差采用无符号运算，可跨越 32 位毫秒计数回绕。",
            "若验证物理时间精度，仪器 GND 接板 GND，探头接 PA0，采集多个边沿；记录相邻边沿间隔与同极性边沿周期，保存仪器采样设置和波形。",
        ],
        "debug": "暂停 CPU 会扰动观测。PC 串口接收时间含缓冲与传输延迟，目标 STATE 中的 ms 也依赖同一个时基；两者只能验证相对行为，不能单独证明 1 ms 绝对精度。",
        "accept": "1 ms 重装配置正确，定时翻转正常；有外部仪器证据才给出物理时间误差，没有则将绝对精度标为未测。",
    },
    {
        "target": "08_Application", "title": "实验八 综合控制与 PC 中文显示",
        "purpose": "用 K1 启停 D1 的一秒闪烁，用 K4 中断触发 PC 显示完成文字。",
        "steps": [
            "选择 08_Application，编译下载。打开 UTF 8 串口终端，复位后 D1 应熄灭。确认启动行为 LAB 8。",
            "短按 K1，串口出现 BLINK ON，D1 立即点亮，此后每 1 s 翻转一次。等待至少三次翻转。",
            "闪烁期间短按 K4，PC 显示“嵌入式实验完成啦”，D1 继续闪烁；长按 K4 应只显示一次，松开再按才新增一行。",
            "再次短按 K1，串口出现 BLINK OFF，D1 立即熄灭并保持。停止闪烁后按 K4，中文仍应正常显示。",
            "重复启停与中文操作至少三轮；同时按下两键时，应处理闪烁状态切换和中文输出两个事件。保存含中文与状态切换的串口日志。",
        ],
        "debug": "K4 对应 EXTI12，和 K1 共用 EXTI15_10_IRQHandler。IRQ 仅锁存事件，中文在主循环发送。Print_Completed 使用明确的 UTF 8 字节串，终端编码需与之匹配。",
        "accept": "两个实际按键分别完成启停与中断中文显示，文字输出期间闪烁继续。若使用调试注入，记录只覆盖被注入后的软件处理链。",
    },
]


def status(value):
    return STATUSES.get(str(value or "not_tested").lower(), "待核实")


def items(value):
    if value is None:
        return []
    return value if isinstance(value, list) else [value]


def text_value(value):
    return str(value) if value is not None else ""


def read_summary(path):
    if not path.exists():
        return {"date": "2026-10-04", "overall": {"status": "not_tested", "conclusion": "八项实验的本次上板验证记录尚待补齐。"}, "labs": []}
    data = json.loads(path.read_text(encoding="utf-8-sig"))
    if not isinstance(data, dict):
        raise ValueError("validation_summary.json must be a JSON object")
    return data


def get_lab_result(summary, number, target, source_path):
    entries = summary.get("labs", [])
    if isinstance(entries, dict):
        entries = [dict(v, target=k) for k, v in entries.items()]
    for result in entries:
        if str(result.get("id", "")) == str(number) or result.get("target") == target:
            result = dict(result)
            source_hash = result.get("source_sha256")
            if source_hash and source_hash.lower() != hashlib.sha256(source_path.read_bytes()).hexdigest():
                result["status"] = "not_tested"
                result["conclusion"] = "当前 main.c 与证据记录的 SHA256 不一致，需对当前版本重新验证。"
                result["checks"] = []
            return result
    return {"status": "not_tested", "conclusion": "本次上板验证待记录。", "checks": []}


def evidence_names(evidence, root):
    output = []
    for entry in items(evidence):
        name = entry.get("path", entry.get("file", "")) if isinstance(entry, dict) else str(entry)
        if not name:
            continue
        try:
            name = str(Path(name).relative_to(root))
        except ValueError:
            try:
                name = str(Path(name).relative_to(root.parent.parent))
            except ValueError:
                pass
        output.append(name)
    return "；".join(output) if output else "尚无证据文件"


def readable_observation(check):
    """Summarize embedded machine records without turning them into new claims."""
    observation = text_value(check.get("observed"))
    try:
        raw = json.loads(observation)
    except (ValueError, TypeError):
        return observation
    if not isinstance(raw, dict):
        return observation
    name = raw.get("name", check.get("item"))
    if name == "coarse_serial_report_alternation":
        values = raw.get("host_intervals_seconds") or []
        average = raw.get("host_mean_seconds")
        text = f"PC 串口到达时间的目标间隔为 {raw.get('expected_seconds')} s。"
        if values and average is not None:
            text += f"采集 {len(values)} 个间隔，平均 {average:.6f} s，最小 {min(values):.6f} s，最大 {max(values):.6f} s。"
        text += "该统计仅为串口粗略时序，不能证明实体灯光、电气波形或校准后的绝对时间精度。"
        return text
    if name == "firmware_reported_1000ms_steps":
        values = raw.get("deltas_ms") or []
        if values:
            return f"相邻 STATE 的目标毫秒计数差共 {len(values)} 项，范围 {min(values)} 至 {max(values)} ms。时间来自同一器件的软件时基，不能独立证明物理 1 ms 精度。"
    if name == "UART_error_counter_not_increasing":
        return f"串口错误计数起始为 {raw.get('first')}，结束为 {raw.get('last')}。"
    if name == "structured_firmware_state_received":
        return "已收到能够解析的 STATE 记录，字段包括毫秒计数、固件报告的 LED 与按键状态以及 IRQ 计数。"
    if name == "real_UART_command_replies_and_ODR_reports":
        return "已执行真实 UART 命令序列并检查回复及固件 ODR 状态报告；命令详情见对应证据 JSON。"
    return observation


class Guide:
    def __init__(self):
        self.doc = Document()
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.top_margin, sec.bottom_margin = Cm(1.85), Cm(1.8)
        sec.left_margin = sec.right_margin = Cm(2)
        sec.footer_distance = Cm(0.85)
        for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
            style = self.doc.styles[name]
            style.font.name = "Calibri"
            style.font.color.rgb = RGBColor(0, 0, 0)
            style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")
        self.doc.styles["Normal"].font.size = Pt(10.5)
        pf = self.doc.styles["Normal"].paragraph_format
        pf.space_after, pf.line_spacing = Pt(5), 1.10
        self.doc.styles["Title"].font.size = Pt(25)
        self.doc.styles["Title"].paragraph_format.space_after = Pt(14)
        self.doc.styles["Subtitle"].font.size = Pt(13)
        self.doc.styles["Heading 1"].font.size = Pt(17)
        self.doc.styles["Heading 1"].paragraph_format.space_after = Pt(9)
        self.doc.styles["Heading 2"].font.size = Pt(12)
        self.doc.styles["Heading 2"].paragraph_format.space_before = Pt(9)
        self.doc.styles["Heading 2"].paragraph_format.space_after = Pt(5)
        code = self.doc.styles.add_style("Source Code", WD_STYLE_TYPE.PARAGRAPH)
        code.font.name, code.font.size = "Consolas", Pt(9)
        code._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "等线")
        code.paragraph_format.space_after = Pt(0)
        code.paragraph_format.line_spacing = 1.05
        code.paragraph_format.keep_together = False
        code.paragraph_format.widow_control = False
        for st in self.doc.styles:
            for border in list(st._element.iter(qn("w:pBdr"))):
                border.getparent().remove(border)
        footer = sec.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.add_run("第 ").font.size = Pt(9)
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), "PAGE")
        footer._p.append(fld)
        footer.add_run(" 页").font.size = Pt(9)
        self.doc.core_properties.title = "普中精灵 F103 上板验证与源码调试流程"
        self.doc.core_properties.subject = "STM32F103C8T6 完整开发板 八项实验"
        self.doc.core_properties.author = ""

    def p(self, text, style=None):
        return self.doc.add_paragraph(str(text), style)

    def h(self, text):
        self.doc.add_heading(text, level=2)

    def page(self, title):
        heading = self.doc.add_heading(title, level=1)
        heading.paragraph_format.page_break_before = True

    def steps(self, values):
        for i, value in enumerate(values, 1):
            p = self.p(f"{i}．{value}")
            p.paragraph_format.space_after = Pt(6)

    def table(self, headers, rows, widths):
        table = self.doc.add_table(rows=1, cols=len(headers))
        table.autofit = False
        for col, width in zip(table.columns, widths):
            col.width = Cm(width)
        for cell, value in zip(table.rows[0].cells, headers):
            cell.text = str(value)
        table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for values in rows:
            for cell, value in zip(table.add_row().cells, values):
                cell.text = str(value)
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            b = OxmlElement("w:" + edge)
            for name, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                b.set(qn("w:" + name), value)
            borders.append(b)
        table._tbl.tblPr.append(borders)
        for rowno, row in enumerate(table.rows):
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            for colno, (cell, width) in enumerate(zip(row.cells, widths)):
                cell.width = Cm(width)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcpr = cell._tc.get_or_add_tcPr()
                margin = OxmlElement("w:tcMar")
                for edge, val in (("top", "75"), ("bottom", "75"), ("left", "95"), ("right", "95")):
                    child = OxmlElement("w:" + edge)
                    child.set(qn("w:w"), val)
                    child.set(qn("w:type"), "dxa")
                    margin.append(child)
                tcpr.append(margin)
                fill = "E7EEF4" if rowno == 0 else "FFFFFF"
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), fill)
                tcpr.append(shading)
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(1)
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.line_spacing = 1.08
                    if width < 2.3:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in p.runs:
                        run.font.size = Pt(9.5)
                        run.bold = rowno == 0
        self.p("").paragraph_format.space_after = Pt(1)
        return table


def add_check_details(g, result, evidence_root):
    checks = items(result.get("checks"))
    evidence_index = {}
    if not checks:
        g.p("本项实测检查尚待记录。按本页步骤完成后，保存现象和证据。")
    for check in checks:
        method = METHODS.get(check.get("method"), text_value(check.get("method")) or "方法待补充")
        name = {
            "coarse_serial_report_alternation": "串口报告交替间隔",
            "firmware_reported_1000ms_steps": "固件报告的 1000 ms 计数差",
            "UART_error_counter_not_increasing": "串口错误计数",
            "structured_firmware_state_received": "结构化状态接收",
            "real_UART_command_replies_and_ODR_reports": "实际串口命令与 ODR 报告",
        }.get(check.get("item"), check.get("item", "检查项"))
        refs = []
        for entry in items(check.get("evidence")):
            path = entry.get("path", entry.get("file", "")) if isinstance(entry, dict) else str(entry)
            if path:
                if path not in evidence_index:
                    evidence_index[path] = len(evidence_index) + 1
                refs.append("E" + str(evidence_index[path]))
        suffix = "  证据 " + " ".join(refs) if refs else ""
        p = g.p(f"{name}：{status(check.get('status'))}（{method}）。{readable_observation(check)}{suffix}")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.05
    # These two statements duplicate the per-lab verdict and visual check,
    # and the common evidence-scope explanation. Retain specific limits.
    repetitive = (
        "整项实验尚未标记完全通过；需结合本表未完成项现场验收。",
        "UART中的LED状态来自固件输出寄存器或意图日志，不能证明LED实体发光。",
    )
    limits = [str(note) for note in items(result.get("limitations"))
              if str(note) not in repetitive]
    if limits:
        g.p("待补测或限制：" + "；".join(note.rstrip("。；") for note in limits) + "。")
    return evidence_index


def make_guide(project, summary_path, output):
    data = read_summary(summary_path)
    evidence_root = Path(data.get("evidence_root", summary_path.parent))
    results = [get_lab_result(data, i, lab["target"], project / "Labs" / lab["target"] / "main.c") for i, lab in enumerate(LABS, 1)]
    overall = data.get("overall", {})
    g = Guide()
    g.p("普中精灵 F103\n上板验证与源码调试流程", "Title")
    g.p("STM32F103C8T6 完整开发板 八项实验", "Subtitle")
    g.p("验证日期  " + data.get("date", "2026-10-04"))
    g.p("工程目录  " + str(project))
    g.p("本手册用于从实验一开始逐项编译、下载和检查板上行为，并附八个实验主程序的完整源码。实验八的中文信息通过底板 CH340 在 PC 串口窗口显示。")
    g.h("本次验证结论")
    g.p("整体状态：" + status(overall.get("status")) + "。" + overall.get("conclusion", "各项上板结果以实际记录为准。"))
    g.p("结果表中的“通过”只覆盖其对应检查项。真实串口、SWD 寄存器读回、调试注入、实体按键和肉眼观察分别标明；调试注入通过不等于实体按键已验证。")
    g.h("底板资源对应关系")
    g.table(["程序资源", "底板资源", "电气含义"], [
        ("LED0", "D1  PA0", "低电平点亮"), ("LED1", "D2  PA1", "低电平点亮"),
        ("KEY0", "K1  PA15", "上拉输入 按下为低"), ("KEY1", "K4  PA12", "上拉输入 按下为低"),
        ("USART1", "PA9 TX  PA10 RX", "经底板 CH340 连接 PC"),
        ("SWD", "PA13 DIO  PA14 CLK", "K3 K2 占用同名引脚 调试时不用"),
    ], [2.8, 5.8, 8.4])
    g.p("D1 D2 指底板一排 LED 中的前两盏。完整底板已有实验所需 LED、按键和串口芯片；RST 只负责复位。")

    g.page("软件环境与上电准备")
    g.table(["软件与组件", "用途与位置"], [
        ("Keil µVision", r"D:\Keil5\UV4\UV4.exe  编辑 编译 SWD 下载和调试"),
        ("Arm Compiler 6", r"D:\Keil5\ARM\ARMCLANG\bin  工程使用 AC6"),
        ("STM32F1 设备包", r"D:\Keil5\Keil\STM32F1xx_DFP\2.4.1  设备 STM32F103C8"),
        ("Tera Term", r"D:\TeraTerm\ttermpro.exe  115200 8N1 无流控 UTF 8"),
        ("FlyMCU", r"D:\FlyMCU\FlyMcu.exe  串口 ISP 下载 HEX"),
        ("CH340 驱动", r"Windows 驱动库  备份 D:\CH340_Driver"),
        ("校验 Python", r"D:\F411_PythonEnv\Scripts\python.exe  目录旧名但检查 F103"),
    ], [4, 13])
    g.p("当前端口记录：" + data.get("serial_port", "待以设备管理器核对") + "。COM 号可能随 USB 端口改变；Tera Term 和 FlyMCU 不能同时占用该串口。")
    g.h("底板跳线与供电")
    g.steps([
        "断电后检查核心板插接方向，移除课设曾接入的外部执行器与模块。J11 短接以启用 LED；点阵旁 J1 接 VCC 侧。",
        "P4 的两组跳帽分别短接 A9–UR 和 A10–UT。普通运行时 BOOT0 和 BOOT1 均接 GND 侧；B1 启动跳帽不是 K1 按键。",
        "USB 数据线连接底板 USB1 即 CH340 USB 转 TTL 接口，打开板上电源开关。核心板 USB 口不是本工程串口，且 PA12 已用作 K4。",
        "打开设备管理器，确认 USB SERIAL CH340 端口无故障标记。打开桌面 STM32实验工具中的 F103 工程和串口入口。",
    ])
    g.h("编译目标与内存")
    g.p("工程为 MDK\\F103_Labs.uvprojx。设备 STM32F103C8，Flash 为 0x08000000 起始的 64 KB，RAM 为 0x20000000 起始的 20 KB。HSE 为 8 MHz，正常 HCLK 为 72 MHz。每次选择对应 Target 后 Rebuild，再使用 Build\\Keil_目标名\\目标名.hex。")

    g.page("两种下载方式与断点调试")
    g.h("ST Link SWD 下载")
    g.steps([
        "断电连接 GND、SWDIO 到 PA13、SWCLK 到 PA14，建议连接 NRST。下载器若有 VTref 或 VAPP 电压参考输入，应按其说明接目标板 3.3 V；先核实针脚功能，不能把标为 3.3 V 的电源输出默认当参考输入。",
        "BOOT0=0。Options for Target → Debug 选择 ST Link Debugger，Settings 选择 SW 和 1 MHz。确认识别到目标后再下载，调试时不要按 K2 K3。",
        "Utilities 使用同一调试器。Flash Download 的 STM32F10x_128.FLM 算法可用于此中容量器件；工程可用容量仍限制为 64 KB。启用 Program、Verify、Reset and Run。",
        "Rebuild 成功后执行 Flash → Download，保存写入及校验成功信息；复位并确认当前 LAB 启动信息。连接失败时检查接线、参考电压和 NRST，必要时使用较低 SWD 速度。底板 USB 供电时不得把另一电源输出直接并联到板上电源。",
    ])
    g.h("串口 ISP 下载")
    g.steps([
        "退出串口终端并释放 COM 口。打开 FlyMCU，选择 STM32 串口下载方式、CH340 实际 COM 口和当前目标 HEX，启用写入后校验。",
        "将 BOOT0 置 1、BOOT1 置 0，复位进入 ROM 下载程序；开始下载，必要时按软件提示再复位。只需下载程序，不修改读保护和选项字节。",
        "校验成功后关闭下载串口，将 BOOT0 恢复 0，再复位运行。打开 Tera Term，复位并保存启动日志。实验一要同时保留 ISP 和 SWD 两份证据。",
    ])
    g.h("调试时的观察方法")
    g.p("Debug → Start Stop Debug Session 进入调试。可在初始化后、主循环事件处理处及 IRQ 入口设断点；F5 继续，F10 单步。Watch 可查看 SystemCoreClock、g_millis、g_key0_irq_count 和 g_key1_irq_count。计时、消抖和串口连续运行的验收必须清除断点后进行。")
    g.p("源码以本工程为准。使用 ST Link 下载或 ISP 都会替换当前应用；切换实验后先核对串口 LAB 编号，避免把前一个程序的现象记到本实验。")

    g.page("脚本下载与串口实测记录")
    g.p("本次实际记录允许采用 ISP 下载加真实串口采集。该路径能继续完成固件写入与串口行为检查，SWD 下载和断点调试的完成情况仍须单列。FlyMCU 和 Keil 的图形界面步骤见前页，以下命令用于复现已保存的自动化记录。")
    g.h("打开 PowerShell 并选定环境")
    g.p("先退出占用 COM 口的 Tera Term、FlyMCU 和其他串口助手。以下示例采用 COM4，重新插拔后先核对实际端口。")
    for command in (
        r"$py = 'D:\F411_PythonEnv\Scripts\python.exe'",
        "Set-Location '" + str(project) + "'",
    ):
        g.p(command, "Source Code")
    g.h("识别器件并保存双读备份")
    g.p("首次改写板上程序前运行下列命令。脚本使用完整普中精灵底板的 CH340 自动复位电路，以 DTR RTS 控制进入 ROM，不要求人工反复切换 BOOT0。只有在本板自动下载电路和跳线已确认可用时，才采用此路径。")
    g.p(r"& $py .\tools\bench_isp.py --port COM4 --identify-backup", "Source Code")
    g.p("该步骤识别 PID、64 KB 容量和 UID，连续读取两次完整 Flash，字节完全一致才保存 double_read_backup。再次运行只会备份运行时的当前程序；后来的备份不能覆盖或取代本次操作前的最早备份。")
    g.h("下载当前实验并全片回读")
    g.p("先在 Keil 中编译对应目标。以实验六为例，执行：")
    g.p(r"& $py .\tools\bench_isp.py --port COM4 --lab 6 --flash", "Source Code")
    g.p("将 6 改成 1 至 8 可选择其他实验。脚本要求同一 UID 已有有效备份，检查 HEX 与编译记录的 SHA256，保存下载前镜像，再按页读改写，最后完整回读 64 KB 与预期镜像逐字节比较。日志中的 PASS 只代表该次下载检查通过。")
    g.h("采集真实串口与执行命令检查")
    g.p(r"& $py .\tools\bench_serial.py --lab 6 --port COM4 --duration 8", "Source Code")
    g.p(r"& $py .\tools\bench_serial.py --lab 5 --port COM4 --duration 90", "Source Code")
    g.p("串口脚本不复位、不下载、不注入 GPIO 或按键。实验六实际发送 Y N 和无效字符；实验一、四、五、八每 200 ms 发送 ? 查询；实验二、三、七只监听。按键实验需在采集窗口内手动操作 K1 K4，并逐条记录动作顺序。")
    g.p("串口 JSON、原始字节和带时间戳的日志保存在 Evidence\\20261004。实验二、三输出的是动作意图日志；STATE 的 LED 值是固件读取 ODR 的报告。两者都不是独立引脚电压或灯光测量，不能据此自动宣布八项全部验收通过。")

    g.page("原有固件的备份与恢复")
    backup = data.get("backup", {})
    original = backup.get("original", {}) if isinstance(backup, dict) else {}
    g.p("原有固件指本次实验修改前读出的板上内容，不能据此推定为出厂固件。最早一次成功的 ISP double_read_backup 才对应本次操作前的状态；后续 preflash 文件可能已经是某个实验程序。")
    g.h("原始备份记录")
    if original:
        binary = original.get("binary", original.get("path", ""))
        g.p("BIN 文件：" + text_value(binary))
        g.p("元数据：" + text_value(original.get("metadata", "见 BIN 同名 JSON")))
        if original.get("sha256"):
            g.p("SHA256", "Normal")
            g.p(str(original["sha256"]), "Source Code")
        g.p("完整长度为 65536 字节，起始地址 0x08000000。UID、读回一致性和磁盘文件校验以原始元数据及 ISP 证据为准。")
    else:
        g.p("在 Backup\\ISP_目标UID\\ 下查找时间最早且有成功 ISP 双读证据关联的 double_read_backup_08000000.bin，并核对同名 JSON。最终恢复前必须明确文件、UID 和 SHA256；不要选择最新 preflash 作为原有固件。")
    for note in items(backup.get("notes") if isinstance(backup, dict) else None):
        g.p(str(note))
    g.h("需要恢复时的操作")
    g.steps([
        "先核对备份元数据中的 UID 与待恢复板一致，确认长度为 64 KB，重新计算文件 SHA256 并与元数据比对。恢复会替换当前实验程序。",
        "备份文件是 BIN。若下载工具支持原始 BIN，明确设置起始地址为 0x08000000；若使用只接收 HEX 的流程，可用 intelhex.bin2hex 以 offset=0x08000000 转换，不能直接把 BIN 改扩展名。",
        "通过已验证可用的 ISP 下载方式，或已经确认连通的 ST Link，写入所选原始镜像并启用校验。当前 bench_isp.py 的 --flash 只接受实验编号，没有通用 --restore 参数。",
        "恢复后完整读取 0x08000000 起的 65536 字节，与所选 BIN 逐字节或以 SHA256 比较；恢复正常 Flash 启动并复位，再检查原程序行为。",
    ])
    g.p("下例仅演示 BIN 转 HEX，先把占位路径换成已核对的最早备份和新输出文件；该命令只转换本地文件，不访问开发板。")
    g.p('& $py -c "from intelhex import bin2hex; bin2hex(r\'完整备份BIN路径\', r\'新输出HEX路径\', offset=0x08000000)"', "Source Code")
    g.p("备份覆盖用户 Flash，不包含选项字节、外部存储器或运行时 RAM；本流程不改变读保护和选项字节。")

    g.page("实际验证结果与证据范围")
    g.table(["实验", "目标", "记录状态", "本次结论"], [
        (str(i), lab["target"], status(result.get("status")), result.get("conclusion", "待补充实测记录。"))
        for i, (lab, result) in enumerate(zip(LABS, results), 1)
    ], [1.2, 3.6, 2.2, 10])
    g.h("实验一下载记录")
    downloads = [entry for entry in items(data.get("downloads"))
                 if entry.get("target") == "01_Platform"
                 or (str(entry.get("method", "")).upper() == "SWD" and not entry.get("target"))]
    if downloads:
        for download in downloads:
            g.p(f"{download.get('method', '下载方式待补充')}：{status(download.get('status'))}。{download.get('notes', '')}")
            g.p("证据：" + evidence_names(download.get("evidence"), evidence_root))
    else:
        g.p("串口 ISP 与 ST Link SWD 的本次实际下载记录均待补充。")
    g.h("尚需补测的内容")
    limitations = items(overall.get("limitations"))
    if limitations:
        for note in limitations:
            g.p(str(note))
    else:
        g.p("具体限制见每项实验的检查记录。没有外部仪器波形时，不给出晶振和 1 ms 节拍的绝对精度结论。")
    g.p("证据根目录：" + str(evidence_root))
    g.p("汇总文件：validation_summary.json。各实验检查方法、读回值和对应日志见后续章节；在复制工程后，应同时保留 Evidence 目录。")

    all_evidence_indexes = []
    for number, (lab, result) in enumerate(zip(LABS, results), 1):
        if number == 1:
            g.page(lab["title"])
        else:
            heading = g.doc.add_heading(lab["title"], level=1)
            heading.paragraph_format.space_before = Pt(16)
        g.p("Target  " + lab["target"] + "\n源码  Labs\\" + lab["target"] + "\\main.c")
        g.p(lab["purpose"])
        g.h("验证调试步骤")
        g.steps(lab["steps"])
        g.h("重点观察")
        g.p(lab["debug"])
        g.h("通过标准")
        g.p(lab["accept"])
        g.h("本次实际结果")
        g.p(status(result.get("status")) + "。" + result.get("conclusion", "本次上板验证待记录。"))
        index = add_check_details(g, result, evidence_root)
        all_evidence_indexes.append((number, lab, index))

    g.page("实测证据文件索引")
    g.p("各实验的 E 编号在本实验内独立编号。下列条目列出与前文检查项对应的文件；以 Build 或 Backup 开头的路径相对于工程根目录，其余文件名相对于 Evidence\\20261004。完整结果和原始日志随工程交付。")
    for number, lab, index in all_evidence_indexes:
        g.h(f"实验 {number}  {lab['target'].replace('_', ' ')}")
        if not index:
            g.p("暂无证据文件。")
            continue
        for path, item_index in index.items():
            p = g.p(f"E{item_index}  " + evidence_names([path], evidence_root))
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.05
            for run in p.runs:
                run.font.size = Pt(9)

    for number, lab in enumerate(LABS, 1):
        source_path = project / "Labs" / lab["target"] / "main.c"
        source_bytes = source_path.read_bytes()
        source = source_bytes.decode("utf-8-sig")
        g.page(f"源码附录 {number}  {lab['title'].split(' ', 1)[1]}")
        g.p("文件  Labs\\" + lab["target"] + "\\main.c")
        g.p("以下为生成本文时该文件的完整内容，可在工程内直接定位；未省略任何主程序语句。")
        for line in source.splitlines():
            g.p(line, "Source Code")
        g.h("文件校验与依赖")
        p = g.p("SHA256  " + hashlib.sha256(source_bytes).hexdigest())
        for run in p.runs:
            run.font.name, run.font.size = "Consolas", Pt(8)
        g.p("公共函数和 IRQ 实现在 App\\platform.c，接口在 App\\platform.h；系统时钟在 App\\system_lab.c。完整源代码、启动文件、标准库和 Keil 工程随源码包提供，运行本 main.c 需要这些依赖。")
        g.p("将本附录与对应实验步骤一起使用。修改源码后须重新编译和复测；旧版本的校验结果不自动适用于修改后的固件。")

    output.parent.mkdir(parents=True, exist_ok=True)
    g.doc.save(output)
    print(json.dumps({"output": str(output), "summary_present": summary_path.exists(), "labs": 8,
                      "source_appendices": 8, "render_required": True}, ensure_ascii=False))


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--project", type=Path, default=PROJECT)
    ap.add_argument("--summary", type=Path, default=SUMMARY)
    ap.add_argument("--output", type=Path, default=OUTPUT)
    ap.add_argument("--check-inputs", action="store_true", help="Read inputs and report status without creating a document")
    args = ap.parse_args()
    missing = [str(args.project / "Labs" / lab["target"] / "main.c") for lab in LABS
               if not (args.project / "Labs" / lab["target"] / "main.c").exists()]
    if missing:
        raise FileNotFoundError("Missing original source files: " + "; ".join(missing))
    if args.check_inputs:
        data = read_summary(args.summary)
        print(json.dumps({"summary_present": args.summary.exists(), "source_files": len(LABS),
                          "overall_status": status(data.get("overall", {}).get("status")),
                          "will_create": str(args.output), "created": False}, ensure_ascii=False))
        return
    make_guide(args.project, args.summary, args.output)


if __name__ == "__main__":
    main()
