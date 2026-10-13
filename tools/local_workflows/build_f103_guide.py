from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path(r'D:\MDK5\F103_Labs'); W=Path(r'D:\Work\嵌入式实验')
d=Document(); sec=d.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(2);sec.bottom_margin=Cm(1.8);sec.left_margin=sec.right_margin=Cm(2)
for name in ('Normal','Title','Subtitle','Heading 1','Heading 2'):
    st=d.styles[name];st.font.name='Calibri';st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'微软雅黑');st.font.color.rgb=RGBColor(0,0,0)
d.styles['Normal'].font.size=Pt(10.5);d.styles['Normal'].paragraph_format.space_after=Pt(7);d.styles['Normal'].paragraph_format.line_spacing=1.15
d.styles['Title'].font.size=Pt(23);d.styles['Heading 1'].font.size=Pt(17);d.styles['Heading 2'].font.size=Pt(12)
for st in d.styles:
    for border in list(st._element.iter(qn('w:pBdr'))):border.getparent().remove(border)
d.core_properties.title='普中精灵 F103 八项实验详细调试流程';d.core_properties.author=''
def p(t,style=None):return d.add_paragraph(t,style)
def h(t):d.add_heading(t,2)
def page(t):d.add_page_break();d.add_heading(t,1)
def steps(items):
    for i,t in enumerate(items,1):p(f'{i}．{t}')
def table(headers,rows,widths=None):
    t=d.add_table(rows=1,cols=len(headers));t.style='Table Grid';t.autofit=False
    for c,v in zip(t.rows[0].cells,headers):c.text=v
    pr=t.rows[0]._tr.get_or_add_trPr();pr.append(OxmlElement('w:tblHeader'))
    for row in rows:
        for c,v in zip(t.add_row().cells,row):c.text=v
    for row in t.rows:
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for i,c in enumerate(row.cells):
            if widths:c.width=Cm(widths[i])
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(5);para.paragraph_format.space_before=Pt(5)
                for run in para.runs:run.font.size=Pt(9.5)
    for c in t.rows[0].cells:
        for run in c.paragraphs[0].runs:run.bold=True
    return t
p('普中精灵 F103\n八项实验详细调试流程','Title')
p('STM32F103C8T6 完整开发板','Subtitle')
p('适用工程  D:\\MDK5\\F103_Labs\n整理日期  2026 年 9 月 29 日\n实验八显示终端  PC 串口窗口')
h('本手册的使用范围')
p('本手册将原实验指导书的八项任务适配到普中精灵完整开发板。先完成第 2 至 4 页的公共准备，再按各实验章节操作。底板已有足够的 LED、按键和 CH340 串口芯片，不必外接第二个 LED、按键或液晶屏。实验一要求的 ST-Link 下载仍需外接 ST-Link；严格测量实验七时间精度时另需示波器或逻辑分析仪。')
h('引脚与实物对应')
table(['程序名称','完整底板资源','配置'],[
('LED0','D1，PA0','低电平亮，高电平灭'),('LED1','D2，PA1','低电平亮，高电平灭'),('KEY0','K1，PA15','上拉输入，按下为低'),('KEY1','K4，PA12','上拉输入，按下为低'),('USART1','PA9 TX，PA10 RX','底板 USB1 → CH340 → PC'),('SWD','PA13 DIO，PA14 CLK','保留调试；K2、K3 不使用')],[3,6,8])
p('这里的 D1、D2 均指底板一排 LED 中的前两个，不是核心板的同名指示灯。RST 是复位键，不能代替实验按键。')
h('验证状态')
p('八个目标已完成 Keil 编译，全部 0 错误、0 警告；工程内存与中断向量检查、八组软件回归通过。F103 新板尚未烧录和实测。各章“预期现象”是验收标准，报告的实际结果应在上板操作后填写。')

page('软件环境与入口')
table(['软件或组件','用途与当前状态'],[
('Keil µVision 5.43.1','已安装：D:\\Keil5\\UV4\\UV4.exe。编辑、编译、SWD 下载与断点调试。'),('Arm Compiler 6.24','已安装并通过本工程编译；Target 中使用 Compiler 6。'),('STM32F1xx DFP 2.4.1','已安装：D:\\Keil5\\Keil\\STM32F1xx_DFP\\2.4.1。设备 STM32F103C8。'),('Tera Term 5.6.2','已安装：D:\\TeraTerm\\ttermpro.exe。115200、8N1、无流控、UTF-8。'),('FlyMCU','已存在：D:\\FlyMCU\\FlyMcu.exe。使用串口 ISP 下载 HEX。'),('CH340 驱动','Windows 驱动库已具备 oem46.inf，3.9.2024.9，微软签名有效；备份在 D:\\CH340_Driver。'),('Python 校验环境','复用 D:\\F411_PythonEnv\\Scripts\\python.exe；旧目录名不影响对 F103 工程的检查。')],[4.4,12.6])
h('打开方式')
steps(['桌面打开“STM32实验工具”，选择“01 F103 Keil工程”。也可直接双击 D:\\MDK5\\F103_Labs\\MDK\\F103_Labs.uvprojx。','选择“02 F103串口终端”时会自动查找 CH340/CH341 COM 口。多个设备时，用 PowerShell 调用 D:\\F103_LabTools\\open_serial.ps1 -Port COMn，COMn 按设备管理器填写。','“03 F103编译与离线检查”会重新编译八个实验并检查软件逻辑，不会下载或擦除开发板。','通过“04 FlyMCU串口下载”打开 ISP 工具。下载和终端不能同时占用同一 COM 口，切换前先断开另一软件。'])
h('安装边界')
p('STM32CubeProgrammer 尚未安装，ST 官网下载未成功。当前工程通过 Keil 加 ST-Link 或 FlyMCU 加 CH340 即可下载，不依赖 CubeProgrammer、CubeMX 或 CubeIDE。Windows 驱动实际驻留系统驱动库，D 盘保存工具和驱动备份。')
p('旧 F411 工程及文档仍保留供追溯。当前实验请使用带 F103 名称的工程和本手册。')

page('上电检查与 Keil 公共设置')
h('断电时先核对底板')
steps(['确认核心板为 STM32F103C8T6，已正确插在普中精灵底板上；拔下课设曾接入的外部执行器与模块，避免引脚冲突。','J11 保持短接，使底板 LED 模块接通；点阵旁 J1 短接到 VCC 侧，禁用点阵输出。不要将 VCC 与 GND 直接短接。','P4 两组跳帽分别短接 A9–UR、A10–UT；正常运行 BOOT0/B0、BOOT1/B1 均接 GND 侧。BOOT 端子的 B1 不是用户按键 K1。','数据线插入底板 USB1（CH340 USB 转 TTL 旁），打开电源开关。核心板自身 USB 口不用于本工程串口；本工程 PA12 用作 K4。','Win+R 输入 devmgmt.msc，在“端口”下记录 USB-SERIAL CH340 的实际 COM 号；没有 COM 口先检查数据线、开关与接口。'])
h('Keil 设置核对')
table(['页面','应有设置'],[
('Device','STM32F103C8；不要选择 F411 或 F103 大容量型号。'),('Target','Cortex-M3，无硬件浮点；IROM1 起始0x08000000、大小0x10000；IRAM1 起始0x20000000、大小0x5000。'),('C/C++','STM32F10X_MD、USE_STDPERIPH_DRIVER、HSE_VALUE=8000000；使用随工程带的 F1 SPL。'),('Output','Create HEX File、Debug Information 已开启。'),('Linker','使用 MDK\\flash.sct；64 KB Flash 与20 KB RAM。'),('时钟','8 MHz HSE ×9，HCLK=72 MHz；SysTick 1 ms；正常串口启动信息应显示 HCLK=72000000。')],[3,14])
p('选择目标后执行 Project → Rebuild all target files。Build Output 出现 0 Error(s), 0 Warning(s) 后，使用 Build\\Keil_目标名\\目标名.hex。不要拿旧 F411 的 HEX 下载到新板。')

page('下载与断点调试通用流程')
h('方式一 ST Link 与 SWD')
steps(['断电连接下载器 GND、SWDIO→PA13、SWCLK→PA14，建议同时连接 NRST。电压参考和供电脚按下载器标识处理；底板 USB 供电时不要再并接另一电源输出。','BOOT0=0。Keil 的 Options for Target → Debug 选择 ST-Link Debugger，Settings 选择 SW、1 MHz。不得按 K2、K3，以免拉低调试引脚。','Utilities 使用同一调试器；Flash Download 保留 STM32F10x_128.FLM 算法，勾选 Program、Verify 与 Reset and Run。算法名称中的128不表示C8具有128 KB，工程仍限制64 KB。','编译通过后点击 Flash → Download。确认下载与校验成功，再观察 LED。若使用普中 CMSIS-DAP 仿真器，需将 Debug/Utilities 改为 CMSIS-DAP、配置SWD；课程点名要求 ST-Link 的部分应使用 ST-Link 完成。'])
h('方式二 FlyMCU 与串口 ISP')
steps(['关闭或断开 Tera Term。FlyMCU 选择 CH340 的 COM 口，选择 STM32 串口下载模式和对应实验 HEX，首次可用115200波特率。应用运行串口与 ROM 下载协议的串口格式由各软件分别管理。','采用手动进入 ROM 的流程：BOOT0=1、BOOT1=0，按 RST，再开始下载。如果软件等待复位，按其提示再次复位；不要依赖未核验的自动 DTR/RTS 下载设置。','下载时启用写入后校验；避免选择解除读保护或修改选项字节。校验成功后停止下载，断电将 BOOT0 恢复0，上电或复位运行。','关闭 FlyMCU 串口，再打开 Tera Term，按 RST 查看本实验启动信息。实验一需要分别保存 ISP 和 ST-Link 两种成功记录。'])
h('断点使用')
p('Debug → Start/Stop Debug Session 进入调试；可在 main 初始化完成后或中断入口设断点，F5继续、F10单步、F11进入函数。在 Watch 中查看 g_millis、g_key0_irq_count、g_key1_irq_count。按键机械抖动可能使原始 IRQ 计数增加多次，业务动作经过20 ms消抖后只执行一次。测时间和串口连续工作时清除断点，全速运行。')

labs=[
('实验一 开发平台的熟悉和使用','01_Platform','D1，底板 CH340；另需 ST-Link 完成第二种下载方式。',[
'选择 01_Platform，Rebuild，确认生成 01_Platform.hex。查看工程组：Experiments 为主程序，Board drivers 为板级驱动，Startup 为中容量启动文件，标准库组为 F1 驱动。',
'先按通用流程使用 FlyMCU ISP 下载，恢复 BOOT0=0并复位；观察 D1 闪烁。保存软件的下载及校验成功画面。',
'再用 ST-Link SWD 下载同一个目标，保存 Keil 下载及校验成功记录。两种方式必须分别实际操作。',
'打开串口终端，复位，确认显示 LAB 1、HCLK=72000000，D1 每500 ms翻转。发送 ? 查看 STATE。'],
['在 main 初始化后设断点，检查 SystemCoreClock=72000000；继续运行再暂停，g_millis 应增加。','阅读启动文件的 Reset_Handler：调用 SystemInit，再进入 C 运行时；platform.c 负责 GPIO、按键、SysTick 和串口初始化。'],
'两种下载均成功且程序运行，能说明启动、时钟、主程序与标准库的关系。只有 ISP 成功时，应将 ST-Link 部分记录为未完成。',
'两种下载成功截图、启动串口记录、板卡接口照片、工程结构及必要库函数说明。'),
('实验二 寄存器点灯','02_Register','底板 D1、D2；无额外 LED 或电阻。',[
'选择 02_Register，编译并下载对应 HEX，复位后观察底板前两个 LED。',
'D1亮D2灭持续约500 ms，随后D1灭D2亮持续约500 ms，反复交替。',
'打开 Labs\\02_Register\\main.c，查看 RCC→APB2ENR 打开GPIOA时钟，以及CRL低8位配置为0x22。',
'观察 GPIOA→BSRR 的两组写值：0x00010002 点亮D1；0x00020001 点亮D2。串口同时输出两灯逻辑状态。'],
['在 Register_LED_Init 完成后看 GPIOA CRL：PA0/PA1为2 MHz推挽输出。','分别停在两次写BSRR后观察ODR低2位：D1亮时为二进制10，D2亮时为01；物理低电平与日志LED=1的含义相反。'],
'两盏实体灯交替亮灭，GPIO配置和控制代码直接访问寄存器；UART和时间辅助函数不改变本实验LED控制方式。',
'关键寄存器代码、CRL/ODR调试截图、两种亮灯状态照片或短视频。'),
('实验三 标准库点灯','03_StdPeriph','底板 D1；无额外器件。',[
'选择 03_StdPeriph，Rebuild后下载，复位开始。',
'观察D1亮1 s、灭1 s，完整亮灭周期约2 s；连续观察至少5个周期。',
'阅读 LED_Init、LED_Write 与 Delay_Ms，理解时钟使能、输出初始化、GPIO_WriteBit和SysTick毫秒计时。',
'打开串口核对 LAB 3，避免误把上一次下载的实验当成当前实验。'],
['在 LED_Init 后观察 GPIOA CRL低8位；在 LED_Write 中查看 index 和 on：on=1会输出Bit_RESET。','不要在 Delay_Ms 的等待循环中停留断点来测一秒；退出调试或清除断点，全速运行后测量。'],
'标准库方式控制D1亮1 s、灭1 s，现象稳定，能解释与实验二寄存器方式的区别。',
'LED初始化与主循环关键代码、运行照片、实测亮灭持续时间。'),
('实验四 按键轮询输入','04_Keys','K1、K4，D1、D2；不使用K2、K3或RST代替。',[
'选择04_Keys，编译下载并复位，两灯初始熄灭。',
'短按K1：两灯同时亮；松开后再按K1：两灯同时灭。长按时只应触发一次。',
'复位后短按K4：D1亮D2灭；再次短按K4：D1灭D2亮。连续重复，验证互斥交替。',
'在互斥状态下按K1，两灯各自反转。两个按键同时有效时，程序先处理K1，再处理K4，因此本轮最终状态由K4的互斥状态决定。',
'快速轻触和长按各测试数次，观察串口STATE输出；每次松开保持超过20 ms再按下。'],
['Keys_Poll 采用20 ms稳定状态过滤；在主循环的 events 判断处设断点，KEY0_EVENT=1、KEY1_EVENT=2。','PA15和PA12未按下为1、按下为0；AFIO_MAPR的SWJ_CFG=010，说明只关闭JTAG而保留SWD。'],
'K1同时翻转两灯，K4独立维持互斥交替；一次按下只执行一次业务动作，无持续连发。',
'按键接线对应表、操作前后LED状态表、消抖代码及STATE日志。'),
('实验五 按键外部中断','05_EXTI','K1→EXTI15→D1，K4→EXTI12→D2。',[
'选择05_EXTI，编译下载并复位，两灯初始熄灭。',
'依次短按K1、K4、K1、K4，预期D1D2状态为10、11、01、00；松开后再按。',
'同时按下K1、K4，检查两灯均能翻转；检查长按只发生一次业务动作。',
'阅读 EXTI15_10_IRQHandler：两个独立if分别处理15和12号线，每条线单独清除挂起位；中断只记录事件，主循环消抖后更新LED。'],
['在 EXTI15_10_IRQHandler 中设置断点，按任一键均可进入；通过 EXTI→PR 或对应IRQ计数识别触发源。','中断计数是原始边沿数，并不保证一次机械按压只加1。验收“一次动作”应看消抖后的LED与STATE次数。'],
'两实体按键均通过外部中断分别控制LED；共享入口不遗漏同时挂起的两个中断。',
'中断配置、共享IRQ处理代码、断点进入截图、四步状态序列与实测日志。'),
('实验六 串口命令控制','06_USART','底板 CH340 与D1；P4需A9–UR、A10–UT短接。',[
'选择06_USART，编译并下载。FlyMCU若仍开着，应先释放串口。',
'打开F103串口入口，确认115200、8位数据、无校验、1位停止、无流控。复位后看到LAB 6和HCLK=72000000。',
'在终端按大写Y，D1应亮；按大写N，D1应灭。终端本地回显关闭，键入字符本身可能不显示。',
'发送小写y、X或回车，LED不应改变；再输入Y、N重复确认。输入?会返回STATE作为附加诊断。',
'若使用其他串口助手，选择文本发送；不要把字符Y写成字符串“0x59”。若使用HEX发送模式，59和4E才分别是Y和N。'],
['USART1 TX为PA9，RX为PA10；可在 USART1_IRQHandler 设断点确认收到字节，再继续查看主循环处理。','正常72 MHz时USART1→BRR为625（0x271）。长时间停在断点可能造成接收溢出，应清除断点后复测。'],
'PC大写Y能点亮D1，大写N能熄灭；其他普通字符不改变LED，串口正常无乱码。',
'COM号与串口参数截图、Y/N操作照片、接收中断和主循环代码、STATE日志。'),
('实验七 SysTick 毫秒定时','07_SysTick','底板D1；精确测量需逻辑分析仪或示波器。',[
'选择07_SysTick，编译下载，打开串口后复位。确认HCLK=72000000，未出现时钟告警。',
'观察D1每1000 ms翻转一次，即亮1 s、灭1 s。串口STATE中的ms应约每1000增加一次。',
'查看 Time_Init 中 SysTick_Config(SystemCoreClock/1000)，SysTick_Handler每次将g_millis加1；主循环采用now-last>=1000并累加last。',
'调试时查看SysTick LOAD=71999、CTRL低3位=111。移除所有断点后再验证闪烁，不用暂停的程序测时间。',
'需要量化误差时，仪器地接板GND、通道接PA0，采集多个边沿；相邻边沿约1 s、同极性边沿约2 s。记录均值、最大偏差及仪器设置。'],
['计时差采用无符号减法，能跨越32位毫秒计数回绕。','串口打印和PC接收存在延迟；STATE与手机秒表只用于初步检查，不能证明1 ms绝对精度。实际精度受晶振、调度和测量设备影响。'],
'正确配置1 ms SysTick节拍并实现1 s翻转；有仪器时另记录物理时间精度，无仪器则将精度测量标为未测。',
'SysTick配置和调度代码、LOAD与SystemCoreClock截图、连续STATE、仪器波形或未测说明。'),
('实验八 综合控制与中文显示','08_Application','K1、K4、D1和PC屏幕；无需OLED或LCD。',[
'选择08_Application，编译下载。打开UTF-8串口终端，复位后D1熄灭。',
'短按K1，终端显示BLINK ON，D1立即亮起，此后每1 s翻转。',
'闪烁过程中短按K4，PC屏幕应显示“嵌入式实验完成啦”；D1继续闪烁。长按K4应只显示一次，松开后再次按下才新增一行。',
'再次短按K1，显示BLINK OFF，D1立即熄灭并保持；停止闪烁时再按K4，中文仍应正常显示。',
'两键同时按下时，也应既处理闪烁状态切换，又处理中文输出。重复启停与输出操作至少3轮。'],
['K4使用EXTI12下降沿中断，经共享EXTI15_10_IRQHandler记录，再由主循环消抖处理；中文发送不在中断内执行。','Print_Completed以明确UTF-8字节发送，终端接收编码应设UTF-8。乱码先查编码；完全无输出先查COM、P4、115200和程序是否全速运行。'],
'两个实体按键分别完成启停闪烁和中断触发中文显示；输出文字时闪烁继续，停止闪烁后中文功能仍可用。',
'包含中文与BLINK ON/OFF的终端截图、两键与D1实物照片、关键事件处理代码、3轮测试记录。')]
for title,target,hw,ss,debug,accept,evidence in labs:
    page(title);p('Keil Target  '+target+'\n源文件  Labs\\'+target+'\\main.c\n硬件  '+hw)
    h('操作步骤');steps(ss)
    h('调试观察');steps(debug)
    h('通过标准');p(accept)
    h('报告应保存的材料');p(evidence)
    p('实际操作记录  日期________  下载方式________  结果________\n异常与处理  __________________________________________')

page('故障排查与结果归档')
table(['现象','检查顺序'],[
('没有CH340串口','底板USB1→数据线→电源开关→设备管理器。驱动已预装；接入后确认设备无黄色感叹号。'),('串口打不开','退出FlyMCU或其他占用程序；多个CH340时明确选择实际COM号。'),('下载后不运行','确认HEX属于F103目标；BOOT0恢复0；复位；查看是否选中了当前实验。'),('D1或D2不亮','核对是底板一排LED，检查J11；J1置VCC侧；低电平才亮。'),('K1没反应','确认按的是K1；程序需释放PA15的JTAG功能；不要用RST或BOOT跳帽代替。'),('SWD连接失败','检查GND、SWDIO、SWCLK、电源参考；不要按K2/K3；使用SW模式，必要时降低速度并连接NRST后选择复位下连接。'),('英文正常中文乱码','接收设置改为UTF-8，重新复位和按K4；不要以GBK解释UTF-8字节。'),('时钟不是72MHz','停止精度验收，检查核心板8MHz晶振与型号。程序显示实际HCLK，HSE失败时可能回退HSI。')],[4,13])
h('提交前核对')
p('每项实验保存当前Target、编译结果、下载校验结果、接线照片、预期与实测对照。报告中明确板型为普中精灵STM32F103C8T6，KEY0/KEY1实际对应K1/K4，实验八在PC显示。旧F411报告中的引脚、时钟和实测结果不能直接套用。')
p('本工程的自动检查结果在 Build\\keil_results.json、project_check_results.json、logic_check_results.json。软件测试执行实际AXF，但外设由模型代替；它不验证电气连接、晶振实频、真实中断响应或烧录成功。')
h('参考资料')
p('1．用户提供的《嵌入式系统实验指导书》。\n2．普中官方《开发板功能及使用介绍》\nhttps://blog.csdn.net/2501_92850798/article/details/149668892\n3．普中官方《独立按键实验》\nhttps://blog.csdn.net/2501_92850798/article/details/150204565\n4．普中官方《LED点阵实验》中的J1与J11说明\nhttps://blog.csdn.net/2501_92850798/article/details/150260376\n5．本机Keil.STM32F1xx_DFP 2.4.1设备定义、ST标准外设库及CMSIS 5.9。')
out=W/'普中精灵F103_八项实验详细调试流程.docx';d.save(out);print(out)
