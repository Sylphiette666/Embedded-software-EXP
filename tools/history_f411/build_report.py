from pathlib import Path
import json, statistics, textwrap
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

WORK=Path(__file__).resolve().parent
PROJECT=Path(r'D:\MDK5\F411_Labs')
OUT=WORK.parent/'嵌入式系统实验报告_NUCLEO_F411RE.docx'
doc=Document()
section=doc.sections[0]
section.page_width=Mm(210);section.page_height=Mm(297)
section.top_margin=Mm(20);section.bottom_margin=Mm(20)
section.left_margin=Mm(22);section.right_margin=Mm(22)
section.header_distance=Mm(8);section.footer_distance=Mm(9)

def font(style,latin,east,size,bold=False):
    style.font.name=latin;style.font.size=Pt(size);style.font.bold=bold
    style.font.color.rgb=RGBColor(0,0,0)
    rf=style.element.get_or_add_rPr().get_or_add_rFonts()
    for k,v in [('ascii',latin),('hAnsi',latin),('eastAsia',east)]:rf.set(qn('w:'+k),v)
normal=doc.styles['Normal'];font(normal,'Times New Roman','宋体',11)
normal.paragraph_format.line_spacing=Pt(14)
normal.paragraph_format.space_after=Pt(4)
normal.paragraph_format.widow_control=True
for name,size in [('Title',24),('Subtitle',13),('Heading 1',16),('Heading 2',12),('Heading 3',11)]:
    style=doc.styles[name];font(style,'Arial','黑体',size,name!='Subtitle')
    style.paragraph_format.space_before=Pt(6 if name.startswith('Heading') else 0)
    style.paragraph_format.space_after=Pt(4)
    style.paragraph_format.keep_with_next=True
    style.paragraph_format.line_spacing=Pt(size+3)
from docx.enum.style import WD_STYLE_TYPE
code_style=doc.styles.add_style('CodeBlock',WD_STYLE_TYPE.PARAGRAPH)
font(code_style,'Consolas','宋体',9)
code_style.paragraph_format.line_spacing=Pt(11)
code_style.paragraph_format.space_after=Pt(0)
code_style.paragraph_format.left_indent=Mm(3)
code_style.paragraph_format.keep_with_next=False
small=doc.styles.add_style('EvidenceNote',WD_STYLE_TYPE.PARAGRAPH)
font(small,'Times New Roman','宋体',9)
small.paragraph_format.line_spacing=1.08;small.paragraph_format.space_after=Pt(5)
caption=doc.styles['Caption'];font(caption,'Times New Roman','宋体',9)
caption.paragraph_format.space_before=Pt(4);caption.paragraph_format.space_after=Pt(4)
caption.paragraph_format.keep_with_next=True
for st in doc.styles:
    if st.type == WD_STYLE_TYPE.PARAGRAPH:
        pp=st.element.get_or_add_pPr()
        grid=OxmlElement('w:snapToGrid');grid.set(qn('w:val'),'0');pp.append(grid)
        for border in list(pp.findall(qn('w:pBdr'))):pp.remove(border)
doc.styles['Subtitle'].font.italic=False

footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=footer.add_run();r.font.size=Pt(9)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
settings=doc.settings.element
update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)

def p(text,style=None):
    para=doc.add_paragraph(text,style)
    if style is None:para.paragraph_format.first_line_indent=Pt(22)
    return para
def heading(text,level=2):return doc.add_heading(text,level)
def chapter(text):
    h=heading(text,1);h.paragraph_format.page_break_before=True;return h
def code(text):
    lines=textwrap.dedent(text).strip('\n').splitlines()
    for i,line in enumerate(lines):
        para=doc.add_paragraph(line,'CodeBlock')
        # Keep small logical chunks together but allow long blocks to paginate.
        para.paragraph_format.keep_with_next=(i<len(lines)-1 and (i+1)%8!=0)
    spacer=doc.add_paragraph();spacer.paragraph_format.space_after=Pt(0);spacer.paragraph_format.line_spacing=Pt(3);spacer.add_run().font.size=Pt(3)
def note(text):return p(text,'EvidenceNote')
def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Mm(w)
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
    pr.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for edge,val in [('top',65),('bottom',65),('left',85),('right',85)]:
        el=OxmlElement('w:'+edge);el.set(qn('w:w'),str(val));el.set(qn('w:type'),'dxa');margins.append(el)
    pr.append(margins)
    for rownum,rowdata in enumerate([headers]+list(rows)):
        row=t.rows[0] if rownum==0 else t.add_row()
        for cell,value,w in zip(row.cells,rowdata,widths):
            cell.width=Mm(w);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            para=cell.paragraphs[0];para.paragraph_format.space_after=Pt(0);para.paragraph_format.line_spacing=Pt(12)
            run=para.add_run(str(value));run.font.size=Pt(9.5);run.bold=(rownum==0)
            if rownum==0:
                shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'E8EDF2');cell._tc.get_or_add_tcPr().append(shading)
        trpr=row._tr.get_or_add_trPr();n=OxmlElement('w:cantSplit');trpr.append(n)
        if rownum==0:
            repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
    spacer=doc.add_paragraph();spacer.paragraph_format.space_after=Pt(0);spacer.paragraph_format.line_spacing=Pt(3);spacer.add_run().font.size=Pt(3)
    return t
def source(name):return (PROJECT/name).read_text(encoding='utf-8')
def between(text,start,end):return text[text.index(start):text.index(end)].rstrip()
platform=source('App/platform.c')
logs={}
for f in (PROJECT/'Build').glob('hardware_*.json'):logs[f.stem]=json.loads(f.read_text(encoding='utf-8'))
def timing(number):
    log=next(v for k,v in logs.items() if k.startswith(f'hardware_{number:02d}_'))
    return next(c['intervals_seconds'] for c in log['checks'] if 'intervals_seconds' in c)
def timing_table(number):
    vals=timing(number)
    table(['有效间隔数','平均值','最小值','最大值'],[[len(vals),f'{statistics.mean(vals):.4f} s',f'{min(vals):.4f} s',f'{max(vals):.4f} s']],[35,43,44,44])
def steps(items):
    for i,item in enumerate(items,1):
        para=p(f'{i}．{item}');para.paragraph_format.first_line_indent=Pt(0)
def common_sections(objective,principle,operations):
    heading('一 实验目的');p(objective)
    heading('二 实验原理');p(principle)
    heading('三 实验步骤');steps(operations)

title=p('嵌入式系统实验报告','Title');title.alignment=WD_ALIGN_PARAGRAPH.CENTER
sub=p('NUCLEO F411RE 平台','Subtitle');sub.alignment=WD_ALIGN_PARAGRAPH.CENTER
meta=p('姓名：____________    学号：________________\n班级：____________    指导教师：____________\n实验日期：2026 年 9 月 29 日')
meta.paragraph_format.first_line_indent=Pt(0);meta.paragraph_format.line_spacing=1.65
heading('实验概况',1)
p('本次实验围绕 GPIO、按键输入、外部中断、串口通信和 SysTick 定时展开，将指导书中的 STM32F103 实现要求迁移到 NUCLEO-F411RE。八个正式实验及一个单板综合演示均完成工程配置和编译，现有硬件条件下完成了下载、GPIO 输出、串口命令、周期闪烁以及单板按键演示的验证。')
p('实际硬件仅有开发板。PA6 的第二个 LED 和 PA0 的第二个按键尚未接入，因此相关实验的完整实物验收仍待补充；实验八暂用电脑串口窗口显示中文。实验一的 UART ISP 下载尚未单独验证。')
table(['实验','代码与构建','本次验证范围'],[
('一 开发平台','通过','ST-Link 下载及运行；ISP 待验证'),
('二 寄存器点灯','通过','PA5 和 PA6 互补输出；缺第二只实物灯'),
('三 标准库点灯','通过','板载 LD2 周期亮灭'),
('四 按键输入','通过','双键双灯程序完成；完整实物验证待补'),
('五 按键中断','通过','两路 EXTI 程序完成；完整实物验证待补'),
('六 串口通信','通过','Y 和 N 控灯及未知字符处理'),
('七 SysTick','通过','重装值与周期功能；精度校准待补'),
('八 综合实验','通过','双键版本待补验收；单板替代演示已验证')],[35,31,100])
note('构建与上板测试记录分别见工程 Build 目录的 keil_results.json、hardware_*.json 和 keil_flash.log。')

chapter('实验环境与硬件配置')
heading('一 开发环境')
table(['项目','实际配置'],[
('开发板与芯片','NUCLEO-F411RE，STM32F411RET6'),
('下载与调试','板载 ST-Link/V2-1，SWD 接口'),
('开发软件','Keil 5.43，Arm Compiler 6.24'),
('器件支持包','STM32F4xx_DFP 2.17.1'),
('驱动与内核文件','STM32F4 SPL 1.8.0，CMSIS Core 5.9'),
('系统时钟','内部 HSI 标称 16 MHz，AHB 与两个 APB 均不分频'),
('串口','USART2，COM3，115200 bit/s，8N1，无流控'),
('工程目录',r'D:\MDK5\F411_Labs'),
('存储器配置','Flash 0x08000000 起，512 KB；SRAM 0x20000000 起，128 KB')],[40,126])
heading('二 引脚分配')
table(['功能','引脚','连接方式与有效状态'],[
('LED0','PA5','板载绿色 LD2，高电平点亮'),
('LED1','PA6 / D12','预留外接 LED，高电平点亮，本次未接'),
('KEY0','PC13','板载蓝色 B1 USER，按下为低电平'),
('KEY1','PA0 / A0','预留按键至 GND，内部上拉，本次未接'),
('串口 TX / RX','PA2 / PA3','USART2 的 AF7，默认连接 ST-Link 虚拟串口'),
('SWD','PA13 / PA14','保留给调试器，不作为实验 GPIO')],[33,36,97])
p('完整实验的补充接线为 PA6 经 330 Ω 电阻接 LED 阳极，LED 阴极接 GND；PA0 经按键接 GND。蓝色 B1 用于用户输入，黑色 RESET 只用于复位。Nucleo 板载 LD2 的点亮电平与指导书中的 F103 mini/精英板不同，程序统一按高电平点亮处理。')
heading('三 数据记录方法')
p('通过串口保存命令应答，通过 SWD 读取 GPIO 输出/输入寄存器并校验 Flash。计时测试使用主机单调时钟和约 10 ms 一次的 SWD 采样，舍弃起始的不完整间隔。下文的间隔均值只用于验证周期功能，不等同于示波器或逻辑分析仪的精度测量。')

chapter('实验一 开发平台的熟悉和使用')
common_sections('熟悉 Nucleo 开发板、MDK 工程结构和编译下载流程，能够根据启动信息及 LED 状态判断程序是否正常运行。',
'复位后，内核从向量表取得栈指针和复位入口，启动文件调用 SystemInit 配置系统环境，再由 C 运行库初始化数据并进入 main。工程需使启动文件、芯片宏、存储器容量和 Flash 算法一致。ST-Link 通过 SWD 下载程序，虚拟串口用于通信，两者用途不同。',[
'创建 F411 工程，添加启动文件、CMSIS、标准库以及应用程序；选择 512 KB Flash 和 128 KB SRAM。',
'编译 01_Platform，使用板载 ST-Link 下载并校验；复位后读取串口启动信息。',
'观察 LD2 周期变化，发送字符 ? 查询状态；整理主要文件和库函数的作用。'])
heading('四 工程结构与关键代码')
table(['文件或模块','作用'],[
('startup_stm32f411xe.s','定义中断向量表和复位入口'),
('system_lab.c','选择 HSI 时钟并设置向量表地址'),
('platform.c 与 platform.h','封装 GPIO、按键、中断、串口和毫秒计时'),
('各实验 main.c','组织对应实验的应用逻辑'),
('flash.sct','规定代码、常量、数据和栈的链接区域')],[67,99])
code('''uint32_t last = Time_Now();
while (1) {
    if ((uint32_t)(Time_Now() - last) >= 500u) {
        last += 500u; LED_Toggle(0);
    }
    if (UART_GetChar() == '?') { Print_State(); }
}''')
heading('五 结果与分析')
p('构建成功，下载后的 Flash 读回内容与目标文件一致。LD2 周期翻转，串口返回状态信息。7 个有效间隔的主机采样平均值为 0.4961 s，与设计的约 0.5 s 翻转相符。工程的九个 Keil Target 均为 0 错误、0 警告。')
p('本次已完成 ST-Link 下载；指导书要求的 UART ISP 下载尚未验证。进入 ROM Bootloader 的启动条件和 UART 引脚必须依据 F411 的资料确认，不能仅以 COM3 存在或串口通信成功代替 ISP 下载验收。')
heading('六 问题与总结')
p('多 Target 工程初期出现多个 main 入口参与链接的问题，修正为每个 Target 只启用对应实验的 main.c。平台搭建应先验证最小程序，再逐项增加外设，这样能把工程配置错误与应用逻辑错误分开定位。')
note('证据：Build/keil_results.json、Build/hardware_01_Platform.json；关键代码：Labs/01_Platform/main.c。')

chapter('实验二 寄存器点灯')
common_sections('掌握通过寄存器配置 GPIO 的方法，实现 LED0、LED1 对应引脚交替输出。',
'F411 的 GPIO 时钟由 AHB1 控制。MODER 每两位决定一个引脚的工作模式，OTYPER 决定推挽或开漏，PUPDR 决定上拉或下拉。BSRR 的低 16 位用于置位，高 16 位用于复位；一次 32 位写入可同时点亮一灯并熄灭另一灯。',[
'使能 GPIOA 时钟，将 PA5、PA6 配置为通用推挽输出，初始置低。',
'在循环中交替写入 BSRR，保持两路输出互补；每次状态维持约 500 ms。',
'下载后读回 GPIOA 输出状态并记录翻转时间。第二个 LED 未接入，实测限于引脚输出。'])
heading('四 关键代码')
code('''#define GPIOA_BSRR32 (*((volatile uint32_t *)(GPIOA_BASE + 0x18u)))
RCC->AHB1ENR |= RCC_AHB1ENR_GPIOAEN;
(void)RCC->AHB1ENR;
GPIOA->BSRRH = (1u << 5) | (1u << 6);
GPIOA->MODER = (GPIOA->MODER & ~((3u << 10) | (3u << 12))) |
                (1u << 10) | (1u << 12);
GPIOA->OTYPER &= ~((1u << 5) | (1u << 6));
GPIOA->OSPEEDR &= ~((3u << 10) | (3u << 12));
GPIOA->PUPDR &= ~((3u << 10) | (3u << 12));''')
code('''GPIOA_BSRR32 = (1u << 5) | (1u << (6 + 16));
Delay_Ms(500);
GPIOA_BSRR32 = (1u << 6) | (1u << (5 + 16));
Delay_Ms(500);''')
note('上面为初始化与循环控制摘录；完整循环还包含串口状态提示。')
heading('五 结果与分析');timing_table(2)
p('采样记录中 PA5/PA6 为交替的（1，0）和（0，1），说明寄存器配置及输出逻辑正确。平均状态间隔为 0.4977 s。由于没有外接第二个 LED，本次结果证明两路引脚输出互补，不能替代两只实物 LED 交替点亮的完整验收。')
heading('六 问题与总结')
p('旧标准库设备头文件将 BSRR 拆为 BSRRL 和 BSRRH 两个半字成员，直接使用 GPIOA->BSRR 会产生编译错误。程序按寄存器地址进行 32 位访问，并用一次写入同时更新两个引脚，避免分次更新时出现短暂的中间状态。')
note('证据：Build/hardware_02_Register.json；关键代码：Labs/02_Register/main.c。')

chapter('实验三 标准库点灯')
common_sections('使用 STM32F4 标准外设库配置 GPIO，完成板载 LED 亮 1 s、灭 1 s，并与寄存器方式比较。',
'标准库将寄存器位操作封装为结构体参数和函数调用。初始化时仍必须先开启 GPIO 时钟，再配置输出模式。库函数提高了可读性，但引脚、有效电平和系统时钟仍需与硬件一致。本板 LD2 为高电平点亮。',[
'通过 RCC_AHB1PeriphClockCmd 开启 GPIOA 时钟，将 PA5 配置为推挽输出。',
'在主循环中交替调用 LED_Write，并使用基于 SysTick 的 Delay_Ms 保持状态。',
'下载程序，检查 PA5 状态及串口 ON/OFF 信息，采集周期数据。'])
heading('四 关键代码')
code('''gpio.GPIO_Pin = LED0_PIN | LED1_PIN;
gpio.GPIO_Mode = GPIO_Mode_OUT;
gpio.GPIO_OType = GPIO_OType_PP;
gpio.GPIO_Speed = GPIO_Speed_2MHz;
gpio.GPIO_PuPd = GPIO_PuPd_NOPULL;
GPIO_Init(GPIOA, &gpio);''')
code('''while (1) {
    LED_Write(0, 1); UART_Puts("ON\\r\\n"); Delay_Ms(1000);
    LED_Write(0, 0); UART_Puts("OFF\\r\\n"); Delay_Ms(1000);
}''')
heading('五 结果与分析');timing_table(3)
p('PA5 按预期交替输出高、低电平，串口对应输出 ON 和 OFF。3 个有效状态间隔的平均值为 0.9926 s；设计的一个完整亮灭周期为 2 s，而不是 1 s。结果符合周期点灯要求，绝对时间偏差的来源包括内部 HSI 容差和主机采样误差。')
heading('六 问题与总结')
p('若沿用 F103 实验板的“低电平点亮”定义，会得到相反现象。本实验把点亮和熄灭封装成 LED_Write，使应用逻辑不直接依赖电平细节。库函数和寄存器方式最终控制的是同一组硬件寄存器，应结合原理理解接口参数。')
note('证据：Build/hardware_03_StdPeriph.json；关键代码：Labs/03_StdPeriph/main.c、App/platform.c。')

chapter('实验四 按键输入')
common_sections('掌握 GPIO 输入和轮询检测，实现一次有效按下对应一次状态变化，并实现双灯翻转及互斥点亮。',
'按键机械触点会在短时间内反复接通和断开，因此不能直接把每次低电平读数都当成一次按下。本程序记录原始电平变化时间，连续稳定 20 ms 后更新稳定状态，只在状态变为按下时产生事件。松开并再次按下后才触发下一次事件。',[
'将 PC13 和 PA0 配置为上拉输入，定义低电平为按下；两灯初始熄灭。',
'循环调用 Keys_Poll 获取消抖后的事件。KEY0 分别反转两灯当前状态；KEY1 切换两个互斥状态。',
'为短按、持续按住、松开及再次按下设计验收用例，待外设齐备后执行完整实物测试。'])
heading('四 关键代码')
code('''uint8_t events = Keys_Poll();
if (events & KEY0_EVENT) { LED_Toggle(0); LED_Toggle(1); }
if (events & KEY1_EVENT) {
    exclusive ^= 1u;
    LED_Write(0, exclusive);
    LED_Write(1, (uint8_t)!exclusive);
}''')
code('''if (down != keys[i].raw) {
    keys[i].raw = down;
    keys[i].changed = now;
}
if (keys[i].stable != down &&
    (uint32_t)(now - keys[i].changed) >= 20u) {
    keys[i].stable = down;
    if (down) { events |= (uint8_t)(1u << i); }
}''')
heading('五 结果与分析')
table(['测试条件','程序规定的结果','当前状态'],[
('KEY0 短按一次','两灯当前状态分别取反','完整实物验证待补'),
('KEY0 持续按住','不重复产生按下事件','正式实验实测待补'),
('KEY1 连续有效按下','在（1，0）与（0，1）之间切换','缺 PA0 按键及 PA6 LED'),
('工程构建','只编译当前 main 入口','0 错误 0 警告')],[41,77,48])
p('上述功能逻辑已完成并通过编译，但本次没有对正式实验四进行完整双键双灯上板验收。单板演示中 B1 的操作记录不能直接替代本实验的 KEY1 实物测试。')
heading('六 问题与总结')
p('按住时不断读取到低电平属于“电平检测”，与“一次按下事件”不同。使用稳定状态和变化时间能够避免按住重复触发，同时不需要在主循环中等待按键松开。使用无符号时间差判断也能处理毫秒计数回绕。')
note('证据：Build/keil_04_Keys.log；关键代码：Labs/04_Keys/main.c、App/platform.c 的 Keys_Poll。')

chapter('实验五 按键中断')
common_sections('理解外部中断来源映射、触发边沿和 NVIC 配置，通过两个按键分别控制两个 LED 状态翻转。',
'F411 使用 SYSCFG 把 GPIO 端口连接到对应编号的 EXTI 线。PC13 对应 EXTI13，由 EXTI15_10_IRQHandler 处理；PA0 对应 EXTI0，由 EXTI0_IRQHandler 处理。两个按键均按下降沿触发。中断负责记录事件和清除标志，主循环完成消抖及灯状态更新。',[
'初始化灯和按键，开启 SYSCFG 时钟，完成 PC13/EXTI13 和 PA0/EXTI0 映射。',
'配置下降沿触发和 NVIC，使能两条中断线；中断发生时清除挂起标志并记录对应事件。',
'主循环调用 Keys_PollIRQ，将消抖后的按下事件与中断记录匹配，再翻转对应 LED。'])
heading('四 关键代码')
code('''SYSCFG_EXTILineConfig(EXTI_PortSourceGPIOC, EXTI_PinSource13);
exti.EXTI_Line = EXTI_Line13;
EXTI_ClearITPendingBit(EXTI_Line13);
EXTI_Init(&exti);
nvic.NVIC_IRQChannel = EXTI15_10_IRQn;
NVIC_Init(&nvic);''')
code(between(platform,'void EXTI15_10_IRQHandler(void)','void EXTI0_IRQHandler(void)'))
code('''uint8_t events = Keys_PollIRQ();
if (events & KEY0_EVENT) { LED_Toggle(0); }
if (events & KEY1_EVENT) { LED_Toggle(1); }''')
heading('五 结果与分析')
p('工程构建为 0 错误、0 警告，两路 EXTI 的初始化和处理函数完整。正式实验五的双键双灯实物验收尚未完成。共享驱动在单板演示中已记录到 B1 触发后产生的功能响应，但该证据只覆盖单板演示的 B1 分支。')
p('g_key0_irq_count 和 g_key1_irq_count 统计的是原始下降沿次数，机械抖动可能使一次按下产生多次中断。因此不能仅用中断次数判断消抖效果，应观察有效事件与 LED 翻转次数是否一致。')
heading('六 问题与总结')
p('移植时需要将 F103 的 AFIO 配置改为 F411 的 SYSCFG。中断中使用长延时会影响其他任务，所以本程序仅记录事件，主循环再处理。读取和清除共享事件位时使用短临界区，避免主循环与中断并发修改造成事件丢失。')
note('证据：Build/keil_05_EXTI.log；关键代码：Labs/05_EXTI/main.c、App/platform.c。')

chapter('实验六 串口通信')
common_sections('掌握串口配置和命令处理，通过电脑发送 Y、N 控制板载 LED，并检查无效输入时的行为。',
'串口通信要求双方的波特率、数据位、校验位和停止位一致。本实验使用 USART2 的 PA2/PA3，并配置 AF7。发送采用查询 TXE，接收采用 RXNE 中断；收到的字节先进入 128 字节环形缓冲区，再由主循环解释命令。DMA 也是串口传输方式之一，但本实验未使用。',[
'初始化 PA2/PA3 复用功能、USART2 时钟和 115200 8N1 参数，使能接收中断。',
'在 COM3 以文本方式依次发送 Y、N、X、Y、N，记录返回信息。',
'通过 SWD 同时读取 PA5 的 ODR 和 IDR，核对输出设置及引脚读回状态。'])
heading('四 关键代码')
code('''GPIO_PinAFConfig(GPIOA, GPIO_PinSource2, GPIO_AF_USART2);
GPIO_PinAFConfig(GPIOA, GPIO_PinSource3, GPIO_AF_USART2);
USART_StructInit(&uart);
uart.USART_BaudRate = 115200u;
uart.USART_Mode = USART_Mode_Rx | USART_Mode_Tx;
USART_Init(USART2, &uart);
USART_ITConfig(USART2, USART_IT_RXNE, ENABLE);
NVIC_EnableIRQ(USART2_IRQn);
USART_Cmd(USART2, ENABLE);''')
code('''int command = UART_GetChar();
if (command == 'Y') {
    LED_Write(0, 1); UART_Puts("OK Y LED0=1\\r\\n");
} else if (command == 'N') {
    LED_Write(0, 0); UART_Puts("OK N LED0=0\\r\\n");
} else if (command == '?') { Print_State(); }''')
heading('五 结果与分析')
uart=logs['hardware_06_USART']
table(['序号','发送字符','串口应答','PA5 ODR','PA5 IDR'],[
(i,c['command'],c['reply'].strip() or '无应答 保持状态',c['PA5_ODR'],c['PA5_IDR']) for i,c in enumerate(uart['checks'],1)], [17,24,75,25,25])
p('五次测试均符合预期。Y 将 PA5 置高，N 将 PA5 置低；X 不改变状态。输出寄存器与引脚读回结果一致，说明命令接收、主循环处理和 GPIO 控制链路正常。')
heading('六 问题与总结')
p('串口助手应使用文本模式发送大写字符；若误用十六进制模式，接收到的字节可能与 Y、N 不同。使用环形缓冲区可以缩短接收中断执行时间。测试时还应检查缓冲区溢出及串口错误计数，不能只依据发送窗口中的字符判断通信成功。')
note('证据：Build/hardware_06_USART.json；关键代码：Labs/06_USART/main.c、App/platform.c。')

chapter('实验七 SysTick 定时')
common_sections('建立标称 1 ms 的硬件计时基准，在此基础上实现 LED 每隔 1 s 翻转，并分析周期测量结果。',
'SysTick 是内核的递减计数器。使用内核时钟时，周期为（LOAD+1）/SystemCoreClock。本工程 SystemCoreClock 为 16 000 000 Hz，SysTick_Config(16000) 对应 LOAD=15999，标称周期为 1 ms。中断只增加毫秒计数，主循环负责周期任务。',[
'初始化 SysTick 并设置中断优先级，确认重装值为 15999。',
'使用无符号时间差判断是否达到 1000 ms；达到后翻转 PA5，基准时间增加 1000 ms。',
'记录串口毫秒计数、GPIO 状态及主机采样间隔，区分标称计时与绝对精度。'])
heading('四 关键代码')
code('''SystemCoreClockUpdate();
if (SysTick_Config(SystemCoreClock / 1000u) != 0u) {
    while (1) {}
}
NVIC_SetPriority(SysTick_IRQn, 3u);
// Interrupt service routine:
void SysTick_Handler(void) { ++g_millis; }''')
code('''uint32_t now = Time_Now();
if ((uint32_t)(now - last) >= 1000u) {
    last += 1000u;
    LED_Toggle(0);
    Print_State();
}''')
heading('五 结果与分析');timing_table(7)
table(['记录项目','第 1 次','第 2 次','第 3 次','第 4 次','第 5 次'],[
('毫秒计数',1011,2011,3011,4011,5011),('LED0 状态',1,0,1,0,1)],[36,26,26,26,26,26])
p('SysTick LOAD 实际读回为 15999，连续串口记录的毫秒计数相差 1000。首个记录从 1011 开始，是因为应用在初始化及输出提示之后才建立计时基准。三个主机采样间隔分别约为 0.9914 s、0.9869 s、0.9941 s，平均 0.9908 s。')
p('相对设计的 1 s，样本均值约偏低 0.92%。该值同时包含 HSI 容差、SWD 采样间隔和主机调度的影响，不能据此直接认定芯片时钟误差。当前已验证周期功能，尚未完成使用外部时钟和测量仪器的精度校准。')
heading('六 问题与总结')
p('相比每次把基准时间设为当前时刻，last += 1000 能减少处理耗时引起的长期相位漂移。硬件计时的分辨率与时钟的绝对准确度不同：配置为 1 ms 计数并不自动保证实测时间精确到 1 ms。')
note('证据：Build/hardware_07_SysTick.json；关键代码：Labs/07_SysTick/main.c、App/platform.c。')

chapter('实验八 综合实验')
common_sections('综合 GPIO、EXTI、USART 和 SysTick，实现按键控制闪烁与中断触发文字显示，并保证两项任务能够同时响应。',
'应用使用 blink_enabled 保存闪烁状态，以毫秒计数安排翻转时刻。按键事件、中断记录、串口输出和定时判断分别处理。避免在主循环中阻塞等待 1 s，可以使闪烁过程中仍能接收按键事件。原题要求两个按键，正式程序使用 B1 和外接 PA0 按键。',[
'初始化 LED、两个按键、USART2 和 SysTick，使能 EXTI13、EXTI0；上电熄灭 LED。',
'KEY0 有效按下时切换 blink_enabled；开启时点亮并重新计时，关闭时立即熄灭。',
'KEY1 通过外部中断产生事件，主循环调用 Print_Completed 发送中文；闪烁功能继续运行。'])
heading('四 正式程序关键代码')
code('''uint8_t events = Keys_PollIRQ();
uint32_t now = Time_Now();
if (events & KEY0_EVENT) {
    blink_enabled ^= 1u;
    last = now;
    LED_Write(0, blink_enabled);
    UART_Puts(blink_enabled ? "BLINK ON\\r\\n" : "BLINK OFF\\r\\n");
}
if (events & KEY1_EVENT) { Print_Completed(); }
if (blink_enabled && (uint32_t)(now - last) >= 1000u) {
    last += 1000u;
    LED_Toggle(0);
}''')
heading('五 结果与分析')
p('正式双键程序完成编译，结果为 0 错误、0 警告。由于 PA0 没有外接按键，原题的第二个实体按键触发显示尚未完成验收。Nucleo 没有板载显示屏，因此当前实现通过电脑串口窗口输出文字；如果课程要求 LCD/OLED，还需要增加屏幕及其驱动。')
heading('六 单板补充验证')
p('针对现有硬件，另建 09_BoardOnly_Demo：短按 B1 切换闪烁，长按至少 800 ms 输出“嵌入式实验完成啦”。B1 下降沿通过 EXTI13 记录，再由主循环消抖、区分短按与长按。此演示用于验证组合功能，不替代原题的双键要求。')
table(['监听经过时间','实际串口记录','对应功能'],[
('4.078 s','嵌入式实验完成啦','长按消息输出'),
('69.031 s','BLINK ON','短按开启闪烁'),
('73.312 s','BLINK OFF','短按关闭闪烁')],[38,75,53])
p('监听过程中电脑只发送状态查询字符 ?，没有发送 B/M 控制命令，因此表中的事件来自实际 B1 操作。串口命令 Y/N、B、M 的辅助测试也通过；这些命令测试与实体按键测试分别保存记录。')
heading('七 问题与总结')
p('综合实验的关键是把功能分成短小的事件处理过程。中断中不执行长延时，主循环按时间差管理 LED，中文以固定 UTF-8 字节发送并由串口助手按 UTF-8 显示。单板长短按方案解决了当前演示条件，但完整课程验收仍需补齐第二个按键。')
note('证据：Build/keil_08_Application.log、hardware_09_BoardOnly_Demo.json、hardware_B1_manual.json；代码：Labs/08_Application/main.c 与 Labs/09_BoardOnly_Demo/main.c。')

chapter('综合分析与实验总结')
heading('一 移植与调试问题')
table(['问题','原因或分析','处理结果'],[
('F103 程序不能直接用于 F411','GPIO 时钟总线、模式寄存器和 EXTI 映射方式不同','改用 AHB1、MODER 和 SYSCFG'),
('CMSIS 与旧标准库不兼容','CMSIS 6 的成员命名与 SPL 的 NVIC->IP 用法不一致','采用 CMSIS 5.9，保留原版驱动'),
('多个 main 重复定义','多 Target 工程中多个入口同时参与构建','每个 Target 显式启用一个入口'),
('调试下载期间发生异常','旧程序的 SysTick/待处理中断可能干扰 RAM Flash 算法','暂停内核并清理中断状态后下载校验'),
('中文显示乱码风险','程序发送编码与串口窗口解码不一致','使用固定 UTF-8 字节并统一解码')],[42,75,49])
heading('二 技术总结')
p('寄存器点灯与标准库点灯体现了同一硬件操作的两个层次。寄存器方式有助于理解位域含义，标准库方式便于组织应用，但两者都必须先核对实际原理图和时钟条件。Nucleo 与原实验板的 LED 电平差异说明，软件移植不能只替换芯片型号。')
p('按键实验需要区分电平、边沿、中断和有效事件。轮询与中断都可能遇到机械抖动，只有对稳定状态进行判定后，才能满足“按一次响应一次”。中断次数和有效操作次数的含义不同，应分别记录。')
p('串口、SysTick 和按键组合后，非阻塞的时间判断能够避免长延时影响输入响应。测试也应逐层展开：先看编译和下载，再看串口应答与寄存器，最后核对真实器件现象。计数值正常只能证明程序逻辑的一部分，绝对计时精度和缺失外设的行为仍需独立测量。')
heading('三 尚需完成的实物验收')
steps(['接入第二个 LED 与按键，完成实验二、四、五、八的全部实物操作并补充照片或视频。',
       '按 F411 Bootloader 的指定引脚和启动方式完成一次 UART ISP 下载，保存下载及运行记录。',
       '明确实验八的显示终端要求；如要求实体 LCD/OLED，则增加硬件及对应显示程序。',
       '需要严格计时精度时，核实外部时钟源，并用示波器或逻辑分析仪测量。'])

chapter('附录 代码与证据索引')
heading('一 工程与日志')
p('工程根目录为 D:\\MDK5\\F411_Labs，Keil 工程为 MDK/F411_Labs.uvprojx。每个实验主程序位于 Labs 下对应的编号目录；公共初始化和中断处理位于 App/platform.c，系统时钟位于 App/system_lab.c。')
table(['记录文件','内容'],[
('Build/keil_results.json','全部九个 Target 的构建结果'),
('Build/keil_flash.log','最终 Keil 下载与 Verify OK 记录'),
('Build/hardware_01_Platform.json','平台测试与 0.5 s 间隔采样'),
('Build/hardware_02_Register.json','互补输出状态及间隔'),
('Build/hardware_03_StdPeriph.json','标准库点灯间隔'),
('Build/hardware_06_USART.json','命令、应答、PA5 ODR 与 IDR'),
('Build/hardware_07_SysTick.json','LOAD、串口毫秒值及间隔'),
('Build/hardware_09_BoardOnly_Demo.json','单板演示串口及定时测试'),
('Build/hardware_B1_manual.json','实际 B1 操作期间的串口记录')],[91,75])
heading('二 下载记录摘录')
code('''Erase Done.
Programming Done.
Verify OK.
Application running ...
Flash Load finished at 14:51:20''')
p('首次修改前已备份完整 512 KB 主 Flash，文件为 Backup/original_flash_08000000.bin，起始地址为 0x08000000。未修改 option bytes；该备份不包含 option bytes 或 OTP。')
heading('三 参考资料')
refs=[
'[1] 张小波．嵌入式系统实验指导书 本科生版．广东工业大学自动化学院物联网与大数据系，2023 年 6 月。',
'[2] STMicroelectronics．UM1724 STM32 Nucleo 64 boards MB1136 用户手册。https://www.st.com/resource/en/user_manual/dm00105823.pdf',
'[3] STMicroelectronics．RM0383 STM32F411xC/E 参考手册。https://www.st.com/resource/en/reference_manual/rm0383-stm32f411xce-advanced-armbased-32bit-mcus-stmicroelectronics.pdf',
'[4] STMicroelectronics．STSW STM32065 STM32F4 标准外设库。https://www.st.com/en/embedded-software/stsw-stm32065.html',
'[5] 本次工程源文件、Keil 构建记录、串口日志及 SWD 测试记录，2026 年 9 月 29 日。']
for ref in refs:p(ref,'EvidenceNote')

doc.core_properties.title='嵌入式系统实验报告'
doc.core_properties.subject='NUCLEO-F411RE GPIO EXTI USART SysTick 实验'
doc.core_properties.author=''
doc.core_properties.keywords='STM32F411RET6 NUCLEO-F411RE 实验报告'
doc.save(OUT)
print(OUT)
print('Paragraphs',len(doc.paragraphs),'Tables',len(doc.tables))
