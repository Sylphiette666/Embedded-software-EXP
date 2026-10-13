from pathlib import Path
from docx import Document
from docx.shared import Mm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

W=Path(__file__).resolve().parent
# Reuse the verified report's typography and table helpers, without executing its content builder.
setup=(W/'build_report.py').read_text(encoding='utf-8').split("title=p('嵌入式系统实验报告'")[0]
exec(compile(setup,str(W/'build_report.py'),'exec'))
OUT=W.parent/'NUCLEO_F411RE_八项实验详细调试流程.docx'
doc.core_properties.title='NUCLEO F411RE 八项实验详细调试流程'
doc.core_properties.subject='软件准备 接线 编译 下载 断点调试 现象验证'
doc.core_properties.author=''

def numbered(items,start=1):
    for i,item in enumerate(items,start):
        para=p(f'{i}．{item}');para.paragraph_format.first_line_indent=Pt(0)
def label(text): p(text,'EvidenceNote')
def lab(title,target,hardware):
    chapter(title)
    p('Keil Target：'+target+'。软件：Keil µVision 5 与 Windows PowerShell 串口窗口。')
    p('硬件：'+hardware)
def finish(criteria,evidence):
    heading('通过标准与记录')
    p(criteria)
    label('保存材料：'+evidence)

t=p('八项实验详细调试流程','Title');t.alignment=WD_ALIGN_PARAGRAPH.CENTER
t=p('NUCLEO F411RE 与 STM32F411RET6','Subtitle');t.alignment=WD_ALIGN_PARAGRAPH.CENTER
p('适用工程：D:\\MDK5\\F411_Labs\n实验八显示终端：PC 串口窗口\n整理日期：2026 年 9 月 29 日','EvidenceNote')
heading('如何使用本手册',1)
p('先完成软件准备、接线和 Keil 公共设置，再按实验编号操作。各章给出了具体 Target、按键顺序、预期现象、断点位置和应保存的材料。文中的预期结果用于指导本次操作，不代表尚未接入的外设已经验收通过。')
p('你使用的是 NUCLEO-F411RE，并非指导书中的正点原子 F103 MINI 板。现有工程已经迁移到 F411；无需重新创建工程，也无需使用 CubeMX 生成另一套程序。')
table(['实验','选择的 Target','额外器件'],[
('一 开发平台','01_Platform','SWD 无；ISP 需 USB 转 TTL 模块及跳线'),
('二 寄存器点灯','02_Register','LED、330 Ω 电阻、连接线'),
('三 标准库点灯','03_StdPeriph','无'),
('四 按键输入','04_Keys','LED、电阻、按键、连接线'),
('五 按键中断','05_EXTI','与实验四共用'),
('六 串口通信','06_USART','无，使用板载虚拟串口'),
('七 SysTick','07_SysTick','功能验证无；精度测量另需仪器'),
('八 综合实验','08_Application','按键、连接线；不需要显示屏')],[34,51,81])
heading('推荐调试顺序')
p('只有开发板时，先做实验一的 SWD 部分，再做三、六、七。接好第二个 LED 和按键后，做二、四、五、八。USB 转 TTL 模块到位后补做实验一的 ISP 下载。器件在各实验中重复使用，不需要每个实验各买一套。')
p('09_BoardOnly_Demo 仅用于单板联调。正式验收实验八时选择 08_Application，使用两个实体按键；B1 长按或串口 M 命令不能代替第二个实体按键。')

chapter('一 软件准备与文件位置')
table(['软件或组件','用途','如何准备'],[
('Keil µVision 5','编译、下载、断点调试','沿用现有 Keil 5.43 和 Arm Compiler 6.24'),
('STM32F4xx_DFP 2.17.1','芯片识别与 Flash 算法','已配置；换电脑时用 Pack Installer 安装同版'),
('STSW-LINK009','ST-Link 与虚拟串口驱动','已有 ST-Link 和 COM 口则继续使用；缺驱动时安装'),
('Windows 设备管理器','查看两个 USB 设备及 COM 号','Windows 自带；Win+R 输入 devmgmt.msc'),
('Windows PowerShell','发送串口字符、接收英文和中文','Windows 自带，配合工程 serial_console.ps1'),
('STM32CubeProgrammer','实验一 UART ISP 烧录','做 ISP 时安装；使用 UART 接口'),
('USB 转 TTL 模块驱动','识别 ISP 用的串口模块','按模块芯片型号安装其厂商驱动'),
('Word 或 WPS','查看流程、整理报告','不参与固件编译与烧录')],[43,49,74])
heading('第一次准备')
numbered([
'保留完整工程目录 D:\\MDK5\\F411_Labs，包含 App、Labs、MDK、Vendor、Build、tools。不要只复制某一个 main.c，也不要在 ZIP 压缩包内直接打开工程。',
'用 USB 数据线连接开发板上 ST-Link 的 USB 接口。保持 CN2 的两个 ST-Link 跳线帽安装，电源配置维持本板原有 USB 供电状态。',
'打开设备管理器，展开“端口（COM 和 LPT）”。记下 STLink Virtual COM Port 的编号：此前为 COM3，换接口后可能变化。',
'打开 Keil，选择 Project → Open Project，打开 D:\\MDK5\\F411_Labs\\MDK\\F411_Labs.uvprojx。后续在同一工程内切换 Target。',
'若缺器件包，打开 Keil 的 Pack Installer，找到 Keil::STM32F4xx_DFP，安装 2.17.1；也可从官方版本页下载 .pack 后安装。保持工程自带的 SPL 1.8.0 和 CMSIS 5.9 文件。'])
heading('软件选择说明')
p('串口显示首选工程附带的 PowerShell 脚本，它已经设置 115200、8N1 和 UTF-8。若使用自己熟悉的图形串口助手，将参数设成相同值即可。Keil 的 Serial Window 不是板载真实 COM 口终端。')
p('指导书提到 FLYMCU，F411 的识别与下载支持需实际核对。本手册给出 ST 官方 STM32CubeProgrammer 的 UART 步骤；若课程指定必须提交 FLYMCU 截图，应先确认老师是否接受 F411 上的软件替换。')
label('官方软件下载入口集中列在最后一章。已有软件可继续使用，无需为本手册重装全部环境。')

chapter('二 接线与供电')
heading('板载资源和实验命名')
table(['实验名称','实际位置','用途'],[
('LED0','绿色 LD2，PA5','程序控制灯，高电平亮'),
('KEY0','蓝色 B1 USER，PC13','用户按键，按下为低电平'),
('RESET','黑色 B2 RESET','重启程序，不作为第二个用户按键'),
('LED1','PA6，Arduino D12','第二个 LED 的外接位置'),
('KEY1','PA0，Arduino A0','第二个按键的外接位置'),
('日常串口','USART2 PA2/PA3','经板载 ST-Link 到电脑，无需外接转换器')],[32,62,72])
heading('第二个 LED 的接线')
numbered(['先拔掉开发板 USB 电源。将 D12 接到 330 Ω 电阻一端，电阻另一端接 LED 阳极；LED 阴极接开发板 GND。',
'新 LED 通常长脚为阳极、短脚及外壳平边一侧为阴极；脚已剪短时以器件标识或万用表二极管档确认。电阻必须串联。'])
code('PA6 / D12 --- 330 ohm --- LED (+)\n                            LED (-) --- GND')
heading('第二个按键的接线')
numbered(['PA0（A0）连接按键的一组触点，按键另一组触点连接 GND。程序内部上拉已经启用，无需另外接上拉电阻。',
'四脚轻触按键有两组内部相连的脚，应选“松开不通、按下才通”的两个触点。不要仅凭脚的相邻位置判断；可用万用表通断档确认。',
'核对面包板导通排与公共 GND，确认没有把 3.3V/5V 直接短接到地，再重新连接 USB。'])
code('PA0 / A0 --- push button --- GND')
heading('接好一次后如何复用')
p('LED 和按键可一直保留，用于实验二、四、五、八。实验三、六、七只操作板载灯，外接器件不会影响正常运行。实验八不需要外接屏幕；不要把串口模块接到 PA2/PA3 与板载串口驱动互相叠加。')
label('只有开发板时，第二灯和第二键的实物验证需等待器件；寄存器数值或电脑命令不能代替实物验收。引脚对应依据 UM1724。')

chapter('三 Keil 公共配置')
p('先在工具栏 Target 下拉框选择本次实验，再打开 Project → Options for Target（Alt+F7）。选项对当前 Target 生效；截图显示的是 09_BoardOnly_Demo，正式实验需切换为对应编号。')
pic=Path(r'C:\Users\35462\AppData\Local\Temp\codex-clipboard-85f0fcd6-ae31-476e-828d-653c5e2498e9.png')
if pic.exists():
    para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.line_spacing=1.0
    rr=para.add_run();rr.add_picture(str(pic),width=Mm(158))
    for dp in para._p.xpath('.//wp:docPr'):dp.set('descr','用户提供的 Keil Target 配置截图，显示 STM32F411RETx 和 Flash RAM 配置。')
    label('图 1 现有工程的 Target 配置页')
table(['页签或项目','核对值'],[
('Device','STMicroelectronics → STM32F411RETx'),
('Target → ARM Compiler','Arm Compiler 6；本机已验证 6.24'),
('Target → MicroLIB / FPU','保留 Use MicroLIB；Single Precision'),
('IROM1','勾选；Start 0x08000000；Size 0x80000；Startup 选中'),
('IRAM1','勾选；Start 0x20000000；Size 0x20000'),
('IROM2 / IRAM2 / 外部存储器','保持未启用；未勾选条目中的数值不参与当前配置'),
('Output','保留 Debug Information 和 Create HEX File'),
('Linker','保留工程的 flash.sct，不改成 F103 的存储器布局')],[53,113])
heading('关于截图里的 Xtal undefined')
p('实际时钟由 App/system_lab.c 配置为内部 HSI 16 MHz，且 SystemCoreClock=16000000。Xtal 显示未定义不能单独说明程序时钟错误；更改这个界面字段也不等于修改 RCC 时钟。先按本工程设置运行，在调试器中核对 SystemCoreClock。')

chapter('四 ST Link 下载与断点调试通用流程')
heading('设置下载器')
numbered(['Options for Target → Debug，选择 Use ST-Link Debugger，不选 Use Simulator。勾选 Load Application at Startup；可保留 Run to main。',
'点击 Debug 页的 Settings，选择当前 ST-Link，接口选 SW/SWD，初次连接使用约 1 MHz 的 SWD 时钟。确认能够识别目标内核。',
'在 Settings 的 Flash Download 页确认有 STM32F4xx 512KB 算法，起始地址 0x08000000、大小 0x80000；勾选 Program、Verify、Reset and Run。常规下载使用 Erase Sectors。',
'Utilities 页选择使用调试驱动进行 Flash 编程（Use Debug Driver），或明确选择 ST-Link。保存设置。不要选 F103 的 Flash 算法。'])
heading('每次更换实验都执行')
numbered(['若处于调试状态，按 Ctrl+F5 退出，再选择本实验 Target；确认左侧树和顶部名称一致。',
'按 F7 编译，检查 Build Output 为 0 Error(s)。首次或修改选项后可用 Project → Rebuild all target files。出现错误时停止下载，避免把上次成功的旧固件当成新程序。',
'点击 Flash → Download（默认 F8），等待 Erase、Programming 和 Verify OK。编译成功与下载成功是两个步骤。',
'打开 PC 串口窗口后，按黑色 RESET 重启。检查启动行中的 LAB 编号，确认板上运行的正是当前实验。',
'先在不设断点的情况下观察完整现象，再进入下述断点调试。每次下载会替换板上当前程序，切换实验时重复以上流程。'])
heading('如何下断点与观察变量')
numbered(['打开当前 Labs/编号目录/main.c，在需要执行的语句行按 F9 设置断点。按 Ctrl+F5 进入调试，按 F5 运行到断点；黄色箭头指向下一条将执行的语句。',
'用 View → Watch Windows → Watch 1 输入变量或寄存器表达式。按 F10 单步越过函数，F11 进入函数，F5 继续运行。局部变量需停在其作用域中观察。',
'结束时删除断点（Ctrl+Shift+F9 删除当前 Target 的所有断点），按 F5 恢复，Ctrl+F5 退出调试，最后按 RESET，重新验证连续运行。'])
p('断点暂停会改变计时和串口接收；不要用暂停时长判断闪烁周期。不要在 Delay_Ms 内持续单步，也不要把 SysTick_Handler 的断点一直开着。变量显示不可用时，先检查作用域；确需逐行观察，可临时把 C/C++ (AC6) 优化设为 -O0，再完整重编译和下载。')
label('首次 Flash 内容备份位于 Backup/original_flash_08000000.bin。常规实验不需要修改 option bytes 或执行解除读保护操作。')

chapter('五 PC 串口窗口的详细操作')
heading('打开已配好的终端')
numbered(['设备管理器中确认 STLink Virtual COM Port 的 COM 号。以下以 COM3 为例；若实际为 COM5，就把命令中的 COM3 改成 COM5。',
'按 Win+R，输入 powershell 并回车，打开交互式 Windows PowerShell 窗口。不要使用 PowerShell ISE，也不要通过输出重定向运行此脚本。',
'逐行执行下面两行命令。第一行保存脚本路径，第二行启动串口终端；ExecutionPolicy Bypass 仅用于这个子进程。'])
code("$console = 'D:\\MDK5\\F411_Labs\\tools\\serial_console.ps1'\npowershell -NoProfile -ExecutionPolicy Bypass -File $console -Port COM3")
numbered(['看到 COM3 115200 8N1 UTF-8 提示后，按开发板黑色 RESET。正常会显示 NUCLEO-F411RE | LAB 编号以及该实验操作说明。',
'需要发命令时让焦点留在串口窗口，按对应字符键即可即时发送，无需按回车。例如英文输入状态下按大写 Y、N 或半角问号 ?。按键可能不会本地回显，以开发板应答为准。',
'完成后按 Esc 关闭串口，释放端口。一个 COM 口同一时间只由一个串口终端使用。下载后若 USB 断开重连，重新打开脚本。'],start=4)
heading('如果你使用图形串口助手')
table(['设置项','数值'],[('端口','实际 STLink Virtual COM Port，例如 COM3'),('波特率','115200'),('数据位 / 停止位','8 / 1'),('校验 / 流控','None / None'),('发送模式','文本或 ASCII；关闭十六进制发送、自动定时发送'),('接收编码','UTF-8，实验八必须能显示中文'),('行结束符','无需自动追加；发送单个字符即可')],[56,110])
heading('命令适用范围')
p('脚本启动提示会列出 Y/N/B/M/?，这是单板演示的命令集合，并非每个实验都支持。实验六支持 Y、N、?；实验八只支持 ? 查询，中文必须由外接按键触发。实验二、三只自动输出状态，不处理 ? 查询。')
p('STATE 中 LED0/LED1 表示灯输出状态；KEY0/KEY1=1 表示“按下”，不是原始高电平。IRQ0/IRQ1 为原始中断次数，可能含抖动。RX_ERRORS 为接收错误或缓冲区溢出计数。')
label('若中文异常，先核对 UTF-8 解码，再检查终端字体；不要为了显示中文而改成 GBK 发送。')

lab('实验一 开发平台与 SWD 下载','01_Platform','开发板、USB 数据线；本页无需外接器件。')
heading('操作流程')
numbered(['打开工程，在下拉框选 01_Platform，打开 Labs/01_Platform/main.c。阅读初始化顺序：Time_Init、LED_Init、Keys_Init、UART_Init。',
'按公共配置核对 STM32F411RETx、Flash/SRAM、ST-Link SWD 和 Flash 算法。按 F7 构建，看到 0 错误后截图。',
'执行 Flash → Download，看到 Verify OK 后截图。若串口尚未打开，按第五章启动 PowerShell 串口终端。',
'按黑色 RESET，核对 LAB 1。观察板载 LD2：每约 0.5 s 改变一次状态，即完整亮灭周期约 1 s。',
'发送半角 ?，应出现一行 STATE，ms 在连续运行时增长。间隔约 1 s 再发一次，确认计数增加。',
'进入 Keil 调试，在 while 前的 uint32_t last = Time_Now(); 行设断点，F5 运行到此处。在 Watch 1 加 SystemCoreClock 和 g_millis。',
'SystemCoreClock 应为 16000000。删除该断点，F5 运行片刻后 Stop，g_millis 应比之前增大；不要期待暂停时还能实时增长。',
'浏览 App/system_lab.c、App/platform.c、Vendor 下的启动文件、MDK/flash.sct，分别记录时钟、中断驱动、复位入口和存储布局的作用。'])
heading('你应该看到什么')
code('NUCLEO-F411RE | LAB 1 | Platform and download smoke test\nLD2 toggles every 500 ms. Send ? for state.\nSTATE ms=... LED0=... LED1=0 KEY0=... KEY1=... IRQ0=0 IRQ1=0 ...')
label('省略号是解释用占位，不是固件输出的固定字符；毫秒值和灯状态随发送时刻变化。')
heading('出现问题时')
p('编译报多个 main：确认只启用当前 Target 对应入口，不要把八个 main.c 全部勾选参与构建。下载后仍显示 LAB 9：当前板上仍是旧演示程序，检查 Target、下载日志及是否复位。只出现电源灯但 LD2 不闪：先退出暂停状态，再确认 BOOT0 没有保持高电平。')
finish('SWD 编译、下载、运行和状态查询全部正常，即完成本实验的 ST-Link 部分。指导书还要求 ISP 下载，应继续下一页。','Target 设置、0 错误构建、Verify OK、LAB 1 串口及灯闪烁照片或视频。')

chapter('实验一 UART ISP 下载补充流程')
p('软件：Keil 生成 HEX；STM32CubeProgrammer 通过 UART 下载。硬件：3.3V TTL 逻辑电平的 USB 转串口模块、杜邦线、BOOT0 跳线帽。以下使用 F411 ROM Bootloader 的 USART1 PA9/PA10；日常 COM3 对应的 PA2/PA3 不能直接套用。')
heading('断电后接线')
table(['连接','开发板端'],[
('模块 TXD → MCU RX','PA10，CN10 第 33 脚，亦为 Arduino D2'),
('模块 RXD ← MCU TX','PA9，CN10 第 21 脚，亦为 Arduino D8'),
('模块 GND','开发板 GND；例如 CN7 第 8 脚'),
('模块 VCC / 5V / 3.3V','不接；开发板继续由自己的 ST-Link USB 供电'),
('BOOT0 拉高','CN7 第 5 脚 VDD 与第 7 脚 BOOT0 用跳线帽连接'),
('BOOT1 拉低','PB2（CN10 第 22 脚）接 GND')],[65,101])
label('排针位置按 UM1724 的 MB1136 / NUCLEO-F411RE 核对，以丝印和 1 脚标记计数。模块必须是 TTL 3.3V 逻辑，不是 RS-232；电源选择跳帽未必同时改变 TXD 电平，应确认模块说明。')
heading('连接 下载 恢复启动')
numbered(['在 Keil 编译 01_Platform，确认 Build/Keil_01_Platform/01_Platform.hex 已更新。退出 Keil 调试，并关闭串口终端。',
'断电接好上表线路；首次 ISP 先拔掉 PA6 外接 LED 和 PA0 按键，减少其他 Bootloader 接口上的干扰。连接开发板 USB 及串口模块 USB，按 RESET，使 BOOT0=1、BOOT1=0 在复位时生效。',
'在设备管理器记录 USB 转 TTL 模块的 COM 号。它通常与 ST-Link 的 COM3 不同，可通过插拔模块区分。',
'启动 STM32CubeProgrammer，接口选 UART，端口选模块对应 COM；波特率先用 115200，Parity 选 Even，数据位 8、停止位 1、流控关闭。点击 Connect，确认识别芯片。',
'进入 Erasing & Programming 页面，Browse 选择上面的 01_Platform.hex，启用 Verify programming，执行 Start Programming。HEX 自带地址；若改用 BIN，起始地址必须填写 0x08000000。',
'等待下载与校验成功，保存界面截图。点击 Disconnect。不要进行 Read Unprotect 或随意修改 option bytes；普通写入实验程序不需要这些操作。',
'拔掉电源，移除 BOOT0 与 VDD 的跳线帽，恢复 BOOT0 默认低电平；移除 PB2 接地线及模块信号线。重新上电，必要时按 RESET。',
'重新打开 ST-Link 对应串口（此前 COM3，115200 8N1），确认 LAB 1 和约 0.5 s 翻转。运行验证仍通过板载 USART2，不是 ISP 模块上的 USART1。'])
p('连接超时：先查 BOOT0/BOOT1 后复位、TX/RX 交叉和共地、模块 COM、Even 校验；重新复位后再连接，可降至 57600。此流程依据 AN2606 的 F411 条目与 UM1724，目前尚未在你的外接模块上实测，完成后再勾选 ISP 验收。')

lab('实验二 寄存器控制两灯交替','02_Register','板载 LD2；外接 PA6/D12 → 330 Ω → LED → GND。')
heading('操作流程')
numbered(['断电接好第二灯，重新上电。选 02_Register，打开 Labs/02_Register/main.c，找到 Register_LED_Init 和主循环中的两次 BSRR 写入。',
'F7 编译，Flash → Download，确认 Verify OK。打开串口，按 RESET，启动行应为 LAB 2。',
'观察至少 5 个交替周期：LD2 亮时外接灯灭；约 0.5 s 后 LD2 灭、外接灯亮。串口应交替打印 LED0=1 LED1=0 与 LED0=0 LED1=1。',
'Ctrl+F5 进入调试，在主循环第一条 GPIOA_BSRR32 写入语句设断点，F5 停在执行前。Watch 中输入下面列出的表达式，先确认初始化已完成。',
'按 F10 执行一次置位/复位写入，PA5 应为 1，PA6 为 0；不要把 BSRR 写寄存器的读回值作为输出状态，以 ODR 和 IDR 为准。',
'在第二条 GPIOA_BSRR32 写入处设断点，用 F5 经过延时到达；F10 执行后应变成 PA5=0、PA6=1。',
'移除断点、退出调试并 RESET。拍摄能同时看见两灯的短视频，保留串口交替日志。'])
heading('建议观察的表达式')
table(['表达式','预期结果'],[
('RCC->AHB1ENR & 1','1，GPIOA 时钟开启'),
('(GPIOA->MODER >> 10) & 0xF','5，即 PA5、PA6 均为 01 输出模式'),
('(GPIOA->ODR >> 5) & 3','1 与 2 交替，分别代表两种互补状态'),
('(GPIOA->IDR >> 5) & 3','通常与相应输出状态一致，用于核对引脚读回')],[91,75])
heading('定位第二灯不亮')
p('先看 PA6 输出是否变化。若 ODR 正常，查 LED 极性、电阻串联位置、面包板孔位与公共地；若 ODR 不变，查是否烧录了 LAB 2。两灯总是同时亮通常是接错引脚，或程序仍在运行实验四。')
finish('两只实物灯互补交替，串口与寄存器状态相符。没有第二灯时只能完成引脚输出检查，保留“实物验证待补”。','两灯交替视频、串口两种状态、ODR 两次不同值、编译下载截图。')

lab('实验三 标准库控制单灯亮灭','03_StdPeriph','开发板和 USB 数据线；观察板载 LD2。')
heading('操作流程')
numbered(['选择 03_StdPeriph，打开 Labs/03_StdPeriph/main.c，确认循环执行 LED_Write(0,1)、Delay_Ms(1000)、LED_Write(0,0)、Delay_Ms(1000)。',
'在 App/platform.c 找 LED_Init，检查 GPIOA 时钟开启、PA5 推挽输出；LED_Write 通过 GPIO_WriteBit 控制电平。',
'F7 构建后下载，打开串口并按 RESET，确认 LAB 3。连续观察约 10 s，LD2 应亮约 1 s、灭约 1 s。',
'核对串口 ON 和 OFF 交替出现。注意“每次状态保持 1 s”意味着完整亮灭周期约 2 s，不能将一整轮误记为 1 s。',
'进入调试，在 LED_Write(0,1) 行设断点，F5 停住后按 F11 进入函数。观察 index=0、on=1，单步或运行到 GPIO_WriteBit 执行后，查看 PA5 输出位。',
'在 LED_Write(0,0) 行设断点，用 F5 经过延时，执行后 PA5 应变为 0。不要在 1000 ms 延时循环内反复单步。',
'删除断点并恢复正常运行，拍摄 LD2 完整亮灭过程；对比实验二直接访问寄存器与本实验库函数调用的区别。'])
heading('建议观察的表达式')
table(['表达式或位置','如何解释'],[
('(GPIOA->ODR >> 5) & 1','1 为 LD2 亮，0 为灭'),
('(GPIOA->MODER >> 10) & 3','初始化后为 1，即通用输出'),
('LED_Write 的 index 与 on','只在函数作用域且变量未被优化时观察'),
('g_millis','自由运行时增加，为 Delay_Ms 提供计时')],[81,85])
heading('出现问题时')
p('亮灭电平相反：本板是高电平点亮，不沿用 F103 MINI 的低电平定义。串口输出 ON 但灯不亮：确认看的是 LD2，而不是电源或 ST-Link 状态灯。程序看起来停在 Delay_Ms：删除时基中断断点，按 F5 让 CPU 连续运行。')
finish('LD2 约 1 s 亮、1 s 灭，ON/OFF 与灯一致，能够指出标准库初始化和写 GPIO 的调用位置。','ON/OFF 串口记录、亮灭视频、GPIO_Init 与 GPIO_WriteBit 相关代码截图。')

lab('实验四 双键轮询控制双灯','04_Keys','板载 B1 和 LD2，外接 PA0 按键与 PA6 LED。')
heading('操作流程')
numbered(['确认第二灯和第二键接线，选 04_Keys，F7 编译并下载。打开串口，按 RESET，确认 LAB 4，初始两灯全灭。',
'短按并松开 B1，一次操作后两灯都亮；再按一次，两灯都灭。每次按下和松开都留出约 0.1 s，便于消抖和观察。',
'按住 B1 约 2 s：应只翻转一次，不应持续反复变化。松开后再按，才再次翻转。',
'按 RESET 回到初始状态，短按 PA0 外接键：第一次 LED0 亮、LED1 灭；第二次 LED0 灭、LED1 亮；继续按应交替。',
'在互斥状态下按 B1，两灯各自取反，因此仍是一亮一灭，只是对调。B1 的含义是“分别翻转当前状态”，不是强制两灯同亮或同灭。',
'每次有效按下会打印 STATE；需要额外查看时发送 ?。按住外接键发送 ? 应看到 KEY1=1，松开后发送应为 0。',
'进入调试，在 if (events & KEY0_EVENT) 行设置条件断点 events != 0（Debug → Breakpoints 可设置条件），F5 后按键。观察 events：B1 为 1，PA0 键为 2。',
'按 F10 观察 LED_Toggle / LED_Write 引起的状态改变；随后删除断点、复位，并重新验证长按不重复。调试停顿本身会影响消抖，不计入最终按键验收。'])
heading('复位后按此顺序核对')
table(['操作序列','LED0','LED1'],[
('RESET','灭','灭'),('B1 第一次','亮','亮'),('B1 第二次','灭','灭'),
('PA0 第一次','亮','灭'),('PA0 第二次','灭','亮'),('此时按 B1','亮','灭')],[104,31,31])
heading('调试重点与故障定位')
p('Watch 观察 events、exclusive（主函数作用域）、GPIOC->IDR 的 bit13 与 GPIOA->IDR 的 bit0。原始 IDR 按下为 0，而串口 KEYx 的“按下”标志为 1。按键一直被判按下时，先查四脚按键是否接到了内部常通的同一组脚。')
finish('两键功能符合表格，持续按住只产生一次有效操作，松开后能够再次触发。','初始和各按键状态、长按视频、STATE 日志、消抖代码截图。')

lab('实验五 双键外部中断控制双灯','05_EXTI','与实验四完全相同，保持外接按键和 LED。')
heading('操作流程')
numbered(['选择 05_EXTI，打开 Labs/05_EXTI/main.c，确认调用 Keys_EnableEXTI(KEY0_EVENT | KEY1_EVENT)。编译下载，打开串口后 RESET，确认 LAB 5。',
'初始两灯都灭。短按 B1，只有 LD2 亮；再按 B1，只有 LD2 灭。确认 PA6 外接灯没有跟着变。',
'短按 PA0 外接键，只有外接 LED 翻转；再次按下，它恢复原状态。确认 LD2 不受影响。',
'分别持续按住两键约 2 s，每次只产生一次有效翻转。最终验收分开按两个键，便于对应事件与现象。',
'发送 ?，查看 IRQ0、IRQ1。B1 操作使 IRQ0 增加，PA0 键使 IRQ1 增加；机械抖动可能令原始次数增加不止 1，而灯只应翻转一次。',
'进入调试，在 App/platform.c 的 EXTI15_10_IRQHandler 中设置断点。F5 运行后按 B1，程序应进入此函数。',
'移除上述断点，在 EXTI0_IRQHandler 内设断点，继续运行后按 PA0 键，应进入第二个处理函数。测试完立即删除中断断点。',
'回到正常运行，RESET 后重新完成两键独立翻转测试，避免把暂停过程引起的额外边沿作为正式结果。'])
heading('中断与事件对应关系')
table(['输入','中断位置','有效功能'],[
('B1 / PC13','EXTI13 → EXTI15_10_IRQHandler','翻转 LED0 / PA5'),
('外接键 / PA0','EXTI0 → EXTI0_IRQHandler','翻转 LED1 / PA6')],[38,85,43])
p('Watch 可添加 g_key0_irq_count、g_key1_irq_count、irq_latched（文件作用域）及主函数中的 events。EXTI->IMR 的 bit0、bit13 应开启；FTSR 对应位为 1，表示下降沿触发。SYSCFG 将 EXTI0 映射到 A 口、EXTI13 映射到 C 口。')
heading('中断没有进入时')
p('先查按下时 IDR 是否从 1 变为 0，再查 SYSCFG 时钟、端口映射、EXTI 屏蔽和 NVIC 使能。程序中清除挂起位后才记录事件，因此在断点处读到 PR 已清零不一定是故障。中断内只记录事件，20 ms 消抖及灯控制在主循环完成。')
finish('两路按键分别进入对应处理函数，并各自只改变对应 LED；长按不连发。','两路中断断点截图、IRQ 计数变化、两灯独立响应视频。')

lab('实验六 PC 串口命令控制 LED','06_USART','只有开发板和 USB 数据线，不使用 ISP 转串口模块。')
heading('操作流程')
numbered(['选择 06_USART，打开 Labs/06_USART/main.c，编译并下载。按第五章打开板载 ST-Link 的 COM 口，115200 8N1，按 RESET 确认 LAB 6。',
'切换英文输入，发送大写 Y。PowerShell 脚本中直接按 Y 即发送；图形串口助手中输入 Y 后点击发送。',
'确认 LD2 点亮并收到 OK Y LED0=1。随后发送大写 N，灯灭并收到 OK N LED0=0。',
'依次测试 Y、N、X、Y、N，按下表核对。X 为未知字符，应没有命令应答，也不改变灯的当前状态。',
'再发送 ?，查看 STATE 和 RX_ERRORS。正常低速操作应没有接收错误。不要开启周期自动发送，以免影响逐条判断。',
'进入调试，在 USART2_IRQHandler 函数中设断点，F5 运行并发送 Y，确认接收中断被触发。移除断点并继续运行，让主循环完成命令处理。',
'若要看解析结果，在主循环 int command = UART_GetChar(); 后的 if 行设置条件断点 command >= 0。观察 Y 为 89（0x59），N 为 78（0x4E）；随后执行灯控制。',
'删除断点，RESET 并重新执行完整命令序列，将本次连续运行的应答记录作为验收依据。'])
table(['顺序','发送','预期应答','LD2'],[
('1','Y','OK Y LED0=1','亮'),('2','N','OK N LED0=0','灭'),
('3','X','无应答，保持上一个状态','灭'),('4','Y','OK Y LED0=1','亮'),('5','N','OK N LED0=0','灭')],[18,24,93,31])
heading('无应答时按此顺序检查')
p('先核对启动行 LAB 6，再检查 COM 号、端口占用、115200 8N1、CPU 是否在运行、字符是否为大写。十六进制发送模式下应发 59 或 4E，但本实验统一使用文本模式，避免把字符串“59”当作 Y。')
p('串口中断会把字节放入缓冲区，主循环才解释命令。单步暂停期间连续发送可能产生溢出；这不代表正常运行时必然有错误。')
finish('五条命令结果全部正确；Y/N 的串口应答、PA5 状态和实际灯一致。','带 LAB 6 标识的命令应答截图、灯亮/灭照片、一次 RXNE 中断或 command 数值截图。')

lab('实验七 SysTick 毫秒时基与周期闪烁','07_SysTick','开发板即可验证功能；示波器或逻辑分析仪仅用于精度测量。')
heading('操作流程')
numbered(['选择 07_SysTick，打开 Labs/07_SysTick/main.c 和 App/platform.c 中的 Time_Init。核对 SysTick_Config(SystemCoreClock / 1000u)。',
'编译下载，打开串口并 RESET，确认 LAB 7。暂时不发送 ?，让程序自动产生周期 STATE，避免手工查询混入测量样本。',
'连续记录至少 10 行自动状态。相邻 ms 应相差 1000，LED0 应按 1、0、1、0…交替；具体起始毫秒值不要求为 0 或 1000。',
'观察 LD2 每约 1 s 翻转，完整亮灭周期约 2 s。串口计数的差值证明程序按 1000 个毫秒节拍调度，不能单凭它证明真实时间毫秒级准确。',
'进入调试，在初始化之后停下，Watch 添加 SystemCoreClock、g_millis、SysTick->LOAD。预期分别为 16000000、当前计数、15999。',
'需要确认中断入口时，可临时在 SysTick_Handler 下一个断点，F5 命中一次后立即删除。不要保留每毫秒触发的断点继续测试灯周期。',
'删除断点，退出调试后 RESET，再采集最终日志。记录系统采用内部 HSI 16 MHz，没有在本流程中进行外部参考时钟校准。'])
heading('计算与观察')
code('SysTick period = (LOAD + 1) / SystemCoreClock\n               = 16000 / 16000000 = 0.001 s\nLED interval   = 1000 ticks = nominal 1 s')
heading('如需测量实际精度')
numbered(['将仪器地线接开发板 GND，信号探头接 PA5（Arduino D13），使用适合 3.3V 信号的输入。保持程序自由运行，不设断点。',
'测量相邻上升沿间隔应约 2 s；测量高电平或低电平持续时间应约 1 s。记录至少 10 个间隔、平均值和偏差。',
'本程序没有把每个 1 ms 节拍直接输出到引脚；1 s 波形只能验证累计调度。若老师要求直接测量 1 ms，还需另做测试引脚脉冲程序，并单独说明测试方式。'])
finish('LOAD 和系统时钟配置正确，自动记录间隔为 1000 ms，LED 周期工作。绝对精度结论以仪器实测为准。','LOAD=15999 截图、至少 10 行周期日志、LED 视频；若有仪器则附波形和误差表。')

lab('实验八 双键综合控制与 PC 中文显示','08_Application','板载 B1、LD2，加 PA0 外接按键；无需外接 LED 或显示屏。')
heading('操作流程')
numbered(['断电检查 PA0 按键到 GND 的连接。选择 08_Application，打开 Labs/08_Application/main.c，确认 blink_enabled 与 Print_Completed 两个功能分支。',
'F7 编译并下载。打开板载 COM 口，使用 UTF-8，按 RESET，确认启动行显示 LAB 8，而不是 LAB 9。上电 LD2 应熄灭。',
'短按并松开 B1：串口出现 BLINK ON，LD2 立即点亮，随后每隔约 1 s 翻转。观察至少 3 次翻转。',
'再次按 B1：出现 BLINK OFF，LD2 立即熄灭并停止闪烁。持续观察至少 3 s，确认保持熄灭。',
'在停止闪烁状态下，短按 PA0 外接键。电脑应显示一次“嵌入式实验完成啦”，灯保持熄灭。',
'再次按 B1 开启闪烁，在灯闪烁过程中按 PA0 键。PC 输出中文，灯继续按原节拍闪烁；该键不改变 blink_enabled。',
'按住 PA0 键约 2 s，应仅产生一次有效中文输出；松开约 0.1 s 后再按，应再次输出。不要把终端重复行或自动发送命令当作按键证据。',
'调试时在 EXTI0_IRQHandler 设断点，F5 后按 PA0，应进入中断。删除断点后，可在主循环 if (events & KEY1_EVENT) 行设置条件 events & 2，观察 Print_Completed 调用。',
'完成断点观察后删除全部断点、RESET，再按“开启闪烁 → 输出中文 → 停止闪烁”顺序完整演示并录屏。'])
table(['操作','PC 屏幕','LD2'],[
('第一次 B1','BLINK ON','开始闪烁'),('PA0 按下一次','嵌入式实验完成啦','保持当前闪烁状态'),
('第二次 B1','BLINK OFF','立即熄灭并停止'),('停止时按 PA0','嵌入式实验完成啦','仍熄灭')],[46,68,52])
heading('调试重点与故障定位')
p('观察 blink_enabled、events、g_key1_irq_count。中文乱码先检查 UTF-8；PA0 键无响应先查 GPIOA->IDR 的 bit0，再查 EXTI0。B1 长按出现中文说明很可能仍在运行 LAB 9，应重新选择并下载 LAB 8。')
finish('两个实体按键各司其职，第二键由中断触发，中文出现在 PC；消息输出与闪烁能同时工作。','LAB 8 启动行、BLINK ON/OFF、中文输出、EXTI0 断点，以及包含实际按键操作的演示视频。')

chapter('六 常见故障与恢复顺序')
table(['现象','优先检查与处理'],[
('找不到头文件或库函数','完整保留 App、Vendor、Labs；不要仅复制 main.c；核对当前 Target 和器件包。'),
('编译器不可用','Target 中选择本机已安装的 Arm Compiler 6；本工程已验证 6.24。'),
('没有 ST-Link 或无法连接','检查 USB 数据线、ST-Link 驱动、CN2 两个跳线帽；关闭其他占用调试器的软件。'),
('Flash Download failed','核对 SWD 和 512KB Flash 算法；先重新插拔并复位；必要时降低 SWD 速度，使用 Connect under Reset / Hardware Reset。'),
('下载成功却没有现象','退出断点暂停、按 RESET、核对 LAB 编号、BOOT0 是否恢复默认低；观察绿色 LD2。'),
('COM 口打不开','关闭其他串口助手及上次脚本；设备管理器查实际端口；重新插拔后重开终端。'),
('窗口空白','打开串口后再 RESET；检查 115200 8N1 和 CPU 正在运行。实验二、三自动输出，不处理 ?。'),
('中文乱码','UTF-8 接收、支持中文的终端字体；工程按 UTF-8 固定字节发送。'),
('按键一直为按下','四脚按键接到了同一导通组，或信号被直接接地；检查 PA0 到 GND 的实际通断。'),
('按一次重复翻转','先移除全部断点并复位，再检查消抖、触点接线和信号干扰；原始 IRQ 多次不等于有效事件多次。'),
('变量 unavailable 或单步跳行','局部变量超出作用域或被优化；在有效作用域断点处查看，必要时临时 -O0 后重建下载。'),
('计时与预计不一致','先删除断点；区分每次翻转和完整周期；内部 HSI 有容差，仪器测量才能给出精度。')],[47,119])
heading('恢复到已知可运行状态')
numbered(['退出 Keil 调试、关闭其他烧录器，按 Esc 关闭串口脚本。',
'断电，移除 ISP 的 BOOT0 高电平跳线和临时串口线，保持正常 USB 供电配置。必要时暂时拔掉外接 LED/按键。',
'重新连接 USB，选择 01_Platform，编译并通过 ST-Link 下载；串口确认 LAB 1 和 LD2 闪烁。',
'基础程序恢复后，再断电接回外设，选择需要的实验 Target。不要用批量删除工程文件或反复更改芯片型号解决连接问题。'])

chapter('七 验收材料与文件整理')
heading('每个实验保存什么')
table(['实验','本次需要保存的核心材料'],[
('一','Keil Target、0 错误、SWD Verify OK、启动串口；ISP 连线与校验成功另存'),
('二','两只灯交替视频、两种串口状态、PA5/PA6 ODR 截图'),
('三','ON/OFF 日志、1 s 亮灭视频、标准库关键调用'),
('四','双键操作顺序及灯状态、长按不重复、消抖判断'),
('五','两路中断各一次命中截图、分别控制两灯的视频'),
('六','Y/N/X/Y/N 应答与灯状态、接收中断或字符数值'),
('七','SysTick LOAD、连续 10 行时间记录、可选仪器测量'),
('八','LAB 8、BLINK ON/OFF 与中文、PA0 中断、双键综合演示')],[27,139])
heading('保存步骤')
numbered(['在工程下手动新建 Evidence 文件夹，按 01_Platform、02_Register 等名称建子文件夹。记录日期、Target、外设接线和实际 COM 号。',
'用 Win+Shift+S 截取 Keil 构建、下载、Watch 或串口界面，保存为 PNG。画面应能看清实验编号、关键值和应答，不只截取“成功”二字。',
'串口内容可从终端复制到 UTF-8 文本文件。要保存长日志，可使用带“接收保存”功能的串口助手，并保持相同参数。',
'灯和实体按键操作用手机录制，镜头同时拍到板子和操作；PC 中文显示可同时录屏。先做一遍无断点完整演示，再保存材料。',
'在现有实验报告中填写本次实际结果；尚未接线或未执行的项目继续标明待验收。已有历史 Build/hardware_*.json 可参考，但不要替代本次自己的操作记录。'])
heading('正式实验中的有效串口命令')
table(['Target','命令或输出'],[
('01 / 04 / 05 / 07 / 08','? 查询状态；04/05 按键事件及 07 定时还会自动输出'),
('02 / 03','自动打印状态，不响应命令'),
('06','Y 点亮、N 熄灭、? 查询；其他字符忽略'),
('09 补充演示','Y/N/B/M/?；不用于代替八项正式实验')],[61,105])

chapter('八 官方资料与软件入口')
p('工程操作和预期现象依据当前源代码；硬件引脚、ROM 下载条件及软件菜单依据下列官方资料。界面文字随版本可能略有差异，按功能名称查找。')
refs=[
('Keil 开发软件','沿用本机 µVision 5；换电脑时从官方安装。','https://www.keil.com/download/product/'),
('STM32F4xx 器件包','版本列表中选择 2.17.1。','https://www.keil.arm.com/packs/stm32f4xx_dfp-keil/versions/'),
('ST-Link Windows 驱动','STSW-LINK009。','https://www.st.com/en/development-tools/stsw-link009.html'),
('STM32CubeProgrammer','实验一的 UART ISP 下载软件。','https://www.st.com/en/development-tools/stm32cubeprog.html'),
('CubeProgrammer UART 操作','端口、波特率及校验设置。','https://dev.st.com/stm32cube-docs/prog/2.23.0/en/docs/markup/Uart_Connection_page.html'),
('UM1724 Nucleo 用户手册','查 NUCLEO-F411RE 引脚图、CN7/CN10 和板载 ST-Link。','https://www.st.com/resource/en/user_manual/dm00105823.pdf'),
('AN2606 ROM Bootloader','查 STM32F411xx、Pattern 1 和 USART1 PA9/PA10。','https://www.st.com/resource/en/application_note/an2606-introduction-to-system-memory-boot-mode-on-stm32-mcus-stmicroelectronics.pdf'),
('STM32F411 数据手册','查 PB2 的 BOOT1 功能及电气条件。','https://www.st.com/resource/en/datasheet/stm32f411re.pdf'),
('µVision 调试菜单','运行、单步、断点和调试会话。','https://www.keil.com/support/man/docs/uv4cl/uv4cl_ui_debug.htm'),
('µVision 快捷键','F7 编译、F8 下载等默认快捷键；若被改动则使用菜单。','https://www.keil.com/download/files/uVision%20Keybooard%20Shortcuts%20v5.30.pdf')]
for i,(name,desc,url) in enumerate(refs,1):
    heading(f'{i} {name}',3)
    p(desc,'EvidenceNote');p(url,'EvidenceNote')
heading('本地文件入口')
code('D:\\MDK5\\F411_Labs\\MDK\\F411_Labs.uvprojx\nD:\\MDK5\\F411_Labs\\README.md\nD:\\MDK5\\F411_Labs\\App\\platform.c\nD:\\MDK5\\F411_Labs\\tools\\serial_console.ps1')
label('实验题目来源：张小波《嵌入式系统实验指导书》本科生版，2023 年 6 月。实验八按本次明确要求，使用 PC 屏幕显示中文。')

doc.save(OUT)
print(OUT)
print('paragraphs',len(doc.paragraphs),'tables',len(doc.tables))
